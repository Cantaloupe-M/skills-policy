#!/usr/bin/env python3
"""Tests for Quality Control Pivot Table Analysis task."""
import pytest
from openpyxl import load_workbook

OUTPUT_FILE = "/root/quality_report.xlsx"

PIVOT_SHEETS = [
    ("Fail Rate by Line", "count", None),
    ("Avg Deviation by Line", "average", None),
    ("Inspections by Shift", "count", None),
    ("Line Shift Matrix", "count", "shift"),
]

REQUIRED_COLUMNS = [
    ("PART_ID", lambda h: "part_id" in h or "partid" in h.replace("_", "")),
    ("PART_NAME", lambda h: "part_name" in h or "partname" in h.replace("_", "")),
    ("LINE", lambda h: h == "line"),
    ("INSPECTOR", lambda h: "inspector" in h),
    ("SHIFT", lambda h: "shift" in h),
    ("DEVIATION_MM", lambda h: "deviation" in h),
    ("WEIGHT_ERROR", lambda h: "weight" in h and "error" in h),
    ("QUALITY_GRADE", lambda h: "quality" in h and "grade" in h or h == "quality_grade"),
]


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Line" in sheet_name and "Shift" not in sheet_name.split("by")[-1].strip():
            assert row_field and "line" in row_field.lower(), f"Row field should be LINE, got '{row_field}'"
        elif "Shift" in sheet_name and "Line" not in sheet_name:
            assert row_field and "shift" in row_field.lower(), f"Row field should be SHIFT, got '{row_field}'"
        elif "Line Shift" in sheet_name:
            assert row_field and "line" in row_field.lower(), f"Row field should be LINE, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and col_field in actual_col.lower(), f"Column field should contain '{col_field}', got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_data(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = [dict(zip(headers, row)) for row in rows[1:] if any(cell is not None for cell in row)]
    return data, headers


VALID_LINES = {"Line-A", "Line-B", "Line-C", "Line-D"}
VALID_SHIFTS = {"Morning", "Afternoon", "Night"}
VALID_GRADES = {"A", "B", "C", "N/A"}


class TestSourceDataContent:
    def test_row_count(self, source_data):
        data, _ = source_data
        assert 800 <= len(data) <= 1200, f"Expected ~1000 rows, got {len(data)}"

    def test_line_values(self, source_data):
        data, headers = source_data
        col = next((h for h in headers if h.lower() == "line"), None)
        vals = {row.get(col) for row in data if row.get(col)}
        invalid = vals - VALID_LINES
        assert not invalid, f"Invalid lines: {invalid}"

    def test_quality_grade_values(self, source_data):
        data, headers = source_data
        col = next((h for h in headers if "quality" in h.lower() and "grade" in h.lower()), None)
        if col:
            vals = {row.get(col) for row in data if row.get(col) is not None}
            invalid = vals - VALID_GRADES
            assert not invalid, f"Invalid quality grades: {invalid}"


class TestDataTransformations:
    def test_deviation_is_absolute(self, source_data):
        """DEVIATION_MM should always be >= 0 (absolute value)."""
        data, headers = source_data
        dev_col = next((h for h in headers if "deviation" in h.lower()), None)
        if not dev_col:
            pytest.skip("No deviation column")
        for i, row in enumerate(data[:100]):
            d = row.get(dev_col)
            if d is not None:
                try:
                    assert float(d) >= 0, f"Row {i+2}: deviation {d} should be non-negative"
                except (ValueError, TypeError):
                    pass

    def test_weight_error_is_ratio(self, source_data):
        """WEIGHT_ERROR should be a reasonable ratio (0 to ~1)."""
        data, headers = source_data
        err_col = next((h for h in headers if "weight" in h.lower() and "error" in h.lower()), None)
        if not err_col:
            pytest.skip("No weight_error column")
        for i, row in enumerate(data[:100]):
            e = row.get(err_col)
            if e is not None:
                try:
                    val = float(e)
                    assert 0 <= val <= 1, f"Row {i+2}: weight_error {val} should be in [0,1]"
                except (ValueError, TypeError):
                    pass

    def test_na_grade_when_missing_measurement(self, source_data):
        """Quality grade should be N/A when deviation is missing."""
        data, headers = source_data
        dev_col = next((h for h in headers if "deviation" in h.lower()), None)
        grade_col = next((h for h in headers if "quality" in h.lower() and "grade" in h.lower()), None)
        if not all([dev_col, grade_col]):
            pytest.skip("Missing columns")
        na_count = 0
        for row in data:
            if row.get(dev_col) is None and row.get(grade_col) == "N/A":
                na_count += 1
        # At least some N/A grades should exist (we know ~5% have missing measurements)
        assert na_count > 0, "Expected some N/A quality grades for rows with missing measurements"

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Fail Rate by Line"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
