#!/usr/bin/env python3
"""Tests for Library Circulation Pivot Table Analysis task."""
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

OUTPUT_FILE = "/root/circulation_report.xlsx"
CATALOG_PDF = "/root/book_catalog.pdf"
CIRCULATION_XLSX = "/root/circulation_records.xlsx"

PIVOT_SHEETS = [
    ("Loans by Genre", "count", None),
    ("Avg Duration by Genre", "average", None),
    ("Loans by Borrower Type", "count", None),
    ("Genre Borrower Matrix", "count", "borrower"),
]

REQUIRED_COLUMNS = [
    ("BOOK_ID", lambda h: "book_id" in h or "bookid" in h.replace("_", "")),
    ("TITLE", lambda h: "title" in h),
    ("GENRE", lambda h: "genre" in h),
    ("AUTHOR", lambda h: "author" in h),
    ("YEAR_PUBLISHED", lambda h: "year" in h and "publish" in h),
    ("BORROWER_TYPE", lambda h: "borrower" in h),
    ("LOAN_DATE", lambda h: "loan_date" in h or "loandate" in h.replace("_", "")),
    ("RETURN_DATE", lambda h: "return_date" in h or "returndate" in h.replace("_", "")),
    ("LOAN_DURATION", lambda h: "duration" in h or "loan_duration" in h),
    ("DECADE", lambda h: "decade" in h),
    ("RETURN_STATUS", lambda h: "return" in h and "status" in h),
    ("WEEKDAY_BUCKET", lambda h: "weekday" in h and "bucket" in h),
]

VALID_GENRES = {"Fiction", "Non-Fiction", "Science", "History", "Biography", "Technology", "Art", "Philosophy"}
VALID_BORROWER_TYPES = {"Student", "Faculty", "Staff", "Community"}
VALID_DECADES = {"1980s", "1990s", "2000s", "2010s", "2020s"}
VALID_RETURN_STATUS = {"returned"}
VALID_WEEKDAY_BUCKET = {"weekday", "weekend"}


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Genre" in sheet_name and "Borrower" not in sheet_name.split("by")[-1].strip():
            assert row_field and "genre" in row_field.lower(), f"Row field should be GENRE, got '{row_field}'"
        elif "Borrower Type" in sheet_name:
            assert row_field and "borrower" in row_field.lower(), f"Row field should be BORROWER_TYPE, got '{row_field}'"
        elif "Genre Borrower" in sheet_name:
            assert row_field and "genre" in row_field.lower(), f"Row field should be GENRE, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and col_field in actual_col.lower(), f"Column field should contain '{col_field}', got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_frame(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    raw_headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data_rows = [row for row in rows[1:] if any(cell is not None for cell in row)]
    return pd.DataFrame(data_rows, columns=raw_headers)


def _find_column(columns, *keywords):
    for col in columns:
        normalized = col.lower().replace("_", "").replace(" ", "")
        if all(keyword in normalized for keyword in keywords):
            return col
    raise AssertionError(f"Missing column with keywords {keywords}. Found: {list(columns)}")


def _weekday_bucket(value):
    ts = pd.Timestamp(value)
    return "weekend" if ts.weekday() >= 5 else "weekday"


def build_expected_frame():
    catalog_tables = pd.read_html(CATALOG_PDF)
    assert catalog_tables, "Could not parse book catalog PDF"
    catalog = pd.concat(catalog_tables, ignore_index=True)
    circulation = pd.read_excel(CIRCULATION_XLSX, dtype=object)

    catalog.columns = [str(c).strip() for c in catalog.columns]
    circulation.columns = [str(c).strip() for c in circulation.columns]

    book_id_col = _find_column(catalog.columns, "book", "id")
    title_col = _find_column(catalog.columns, "title")
    genre_col = _find_column(catalog.columns, "genre")
    author_col = _find_column(catalog.columns, "author")
    year_col = _find_column(catalog.columns, "year", "publish")

    circ_book_col = _find_column(circulation.columns, "book", "id")
    borrower_col = _find_column(circulation.columns, "borrower")
    loan_date_col = _find_column(circulation.columns, "loan", "date")
    return_date_col = _find_column(circulation.columns, "return", "date")

    catalog = catalog.rename(
        columns={
            book_id_col: "BOOK_ID",
            title_col: "TITLE",
            genre_col: "GENRE",
            author_col: "AUTHOR",
            year_col: "YEAR_PUBLISHED",
        }
    )
    circulation = circulation.rename(
        columns={
            circ_book_col: "BOOK_ID",
            borrower_col: "BORROWER_TYPE",
            loan_date_col: "LOAN_DATE",
            return_date_col: "RETURN_DATE",
        }
    )

    circulation["BOOK_ID"] = circulation["BOOK_ID"].map(lambda v: None if pd.isna(v) else str(v).strip())
    circulation["LOAN_DATE"] = pd.to_datetime(circulation["LOAN_DATE"], errors="coerce")
    circulation["RETURN_DATE"] = pd.to_datetime(circulation["RETURN_DATE"], errors="coerce")

    merged = circulation.merge(catalog, on="BOOK_ID", how="left")

    rows = []
    for row in merged.to_dict("records"):
        if row["BOOK_ID"] is None or pd.isna(row["TITLE"]):
            continue
        if pd.isna(row["LOAN_DATE"]) or pd.isna(row["RETURN_DATE"]):
            continue
        duration = int((row["RETURN_DATE"] - row["LOAN_DATE"]).days)
        if duration <= 0:
            continue
        year = int(float(row["YEAR_PUBLISHED"]))
        rows.append(
            {
                "BOOK_ID": row["BOOK_ID"],
                "TITLE": row["TITLE"],
                "GENRE": row["GENRE"],
                "AUTHOR": row["AUTHOR"],
                "YEAR_PUBLISHED": year,
                "BORROWER_TYPE": row["BORROWER_TYPE"],
                "LOAN_DATE": row["LOAN_DATE"].strftime("%Y-%m-%d"),
                "RETURN_DATE": row["RETURN_DATE"].strftime("%Y-%m-%d"),
                "LOAN_DURATION": duration,
                "DECADE": f"{(year // 10) * 10}s",
                "RETURN_STATUS": "returned",
                "WEEKDAY_BUCKET": _weekday_bucket(row["LOAN_DATE"]),
            }
        )

    expected = pd.DataFrame(rows)
    expected = expected.drop_duplicates().reset_index(drop=True)
    return expected


@pytest.fixture(scope="module")
def expected_frame():
    return build_expected_frame()


class TestSourceDataContent:
    def test_row_count_matches_expected(self, source_frame, expected_frame):
        assert len(source_frame) == len(expected_frame)

    def test_genre_values(self, source_frame):
        col = _find_column(source_frame.columns, "genre")
        vals = {str(v) for v in source_frame[col].dropna().unique()}
        invalid = vals - VALID_GENRES
        assert not invalid, f"Invalid genres: {invalid}"

    def test_borrower_values(self, source_frame):
        col = _find_column(source_frame.columns, "borrower")
        vals = {str(v) for v in source_frame[col].dropna().unique()}
        invalid = vals - VALID_BORROWER_TYPES
        assert not invalid, f"Invalid borrower types: {invalid}"

    def test_decade_values(self, source_frame):
        col = _find_column(source_frame.columns, "decade")
        vals = {str(v) for v in source_frame[col].dropna().unique()}
        invalid = vals - VALID_DECADES
        assert not invalid, f"Invalid decades: {invalid}"

    def test_extra_metadata(self, source_frame):
        return_status_col = _find_column(source_frame.columns, "return", "status")
        weekday_col = _find_column(source_frame.columns, "weekday", "bucket")
        assert set(source_frame[return_status_col].dropna().unique()).issubset(VALID_RETURN_STATUS)
        assert set(source_frame[weekday_col].dropna().unique()).issubset(VALID_WEEKDAY_BUCKET)


class TestDataTransformations:
    def test_loan_duration_positive(self, source_frame):
        dur_col = _find_column(source_frame.columns, "duration")
        assert (source_frame[dur_col].astype(float) > 0).all()

    def test_loan_duration_range(self, source_frame):
        dur_col = _find_column(source_frame.columns, "duration")
        values = source_frame[dur_col].astype(float)
        assert ((values >= 1) & (values <= 365)).all()

    def test_source_matches_independent_expected(self, source_frame, expected_frame):
        ordered_columns = [
            "BOOK_ID",
            "TITLE",
            "GENRE",
            "AUTHOR",
            "YEAR_PUBLISHED",
            "BORROWER_TYPE",
            "LOAN_DATE",
            "RETURN_DATE",
            "LOAN_DURATION",
            "DECADE",
            "RETURN_STATUS",
            "WEEKDAY_BUCKET",
        ]
        actual = source_frame[ordered_columns].copy()
        expected = expected_frame[ordered_columns].copy()
        for col in ["YEAR_PUBLISHED", "LOAN_DURATION"]:
            actual[col] = actual[col].astype(float).round(6)
            expected[col] = expected[col].astype(float).round(6)
        pd.testing.assert_frame_equal(actual.reset_index(drop=True), expected.reset_index(drop=True), check_dtype=False)

    def test_pivot_aggregates_match_expected(self, source_frame, expected_frame):
        actual_loans_by_genre = source_frame.groupby(_find_column(source_frame.columns, "genre")).size().to_dict()
        expected_loans_by_genre = expected_frame.groupby("GENRE").size().to_dict()
        assert actual_loans_by_genre == expected_loans_by_genre

        actual_duration = source_frame.groupby(_find_column(source_frame.columns, "genre"))[_find_column(source_frame.columns, "duration")].mean().round(6).to_dict()
        expected_duration = expected_frame.groupby("GENRE")["LOAN_DURATION"].mean().round(6).to_dict()
        assert actual_duration == expected_duration

        actual_borrowers = source_frame.groupby(_find_column(source_frame.columns, "borrower")).size().to_dict()
        expected_borrowers = expected_frame.groupby("BORROWER_TYPE").size().to_dict()
        assert actual_borrowers == expected_borrowers

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Loans by Genre"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
