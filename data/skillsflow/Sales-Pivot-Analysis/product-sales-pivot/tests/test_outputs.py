#!/usr/bin/env python3
"""Tests for Product Sales Pivot Table Analysis task."""
import math
from collections import defaultdict

import pytest
from openpyxl import load_workbook
import pandas as pd

OUTPUT_FILE = "/root/product_sales_report.xlsx"
CATALOG_PDF = "/root/product_catalog.pdf"
SALES_XLSX = "/root/sales_transactions.xlsx"

PIVOT_SHEETS = [
    ("Revenue by Category", "sum", None),
    ("Units by Region", "sum", None),
    ("Products by Category", "count", None),
    ("Category Region Matrix", "sum", "region"),
]

REQUIRED_COLUMNS = [
    ("PRODUCT_ID", lambda h: "product_id" in h or "productid" in h.replace("_", "")),
    ("PRODUCT_NAME", lambda h: "product_name" in h or "productname" in h.replace("_", "")),
    ("CATEGORY", lambda h: "category" in h),
    ("REGION", lambda h: "region" in h),
    ("QUANTITY", lambda h: "quantity" in h),
    ("UNIT_PRICE", lambda h: "unit_price" in h or "unitprice" in h.replace("_", "")),
    ("UNIT_COST", lambda h: "unit_cost" in h or "unitcost" in h.replace("_", "")),
    ("REVENUE", lambda h: "revenue" in h),
    ("PROFIT", lambda h: "profit" in h and "margin" not in h),
    ("MARGIN_PCT", lambda h: "margin" in h),
    ("PRICE_STATUS", lambda h: "price" in h and "status" in h),
    ("CATALOG_MATCH_STATUS", lambda h: "catalog" in h and "match" in h),
    ("RECONCILIATION_ACTION", lambda h: "reconciliation" in h and "action" in h),
]

VALID_CATEGORIES = {"Electronics", "Office Supplies", "Furniture", "Software"}
VALID_REGIONS = {"North", "South", "East", "West"}
ALLOWED_PRICE_STATUS = {"catalog_price", "transaction_override"}
ALLOWED_MATCH_STATUS = {"matched", "unmatched"}
ALLOWED_ACTIONS = {
    "none",
    "dropped_missing_product_id",
    "dropped_unknown_product_id",
    "dropped_nonpositive_quantity",
    "trimmed_region",
    "normalized_region_case",
    "normalized_month_case",
    "normalized_quarter_case",
    "deduplicated_exact_duplicate",
    "filled_unit_price_from_catalog",
    "used_transaction_unit_price",
}


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Category" in sheet_name or "category" in sheet_name.lower():
            assert row_field and "category" in row_field.lower(), f"Row field should be CATEGORY, got '{row_field}'"
        elif "Region" in sheet_name:
            assert row_field and "region" in row_field.lower(), f"Row field should be REGION, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and col_field in actual_col.lower(), f"Column field should be '{col_field}', got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_frame(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    raw_headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data_rows = [row for row in rows[1:] if any(cell is not None for cell in row)]
    return pd.DataFrame(data_rows, columns=raw_headers)


def _find_column(columns, *keywords):
    for col in columns:
        normalized = col.lower().replace("_", "").replace(" ", "")
        if all(keyword in normalized for keyword in keywords):
            return col
    raise AssertionError(f"Missing column with keywords {keywords}. Found: {list(columns)}")


def _normalize_key(value):
    return " ".join(str(value).strip().split()).lower()


def _normalize_region(value):
    if pd.isna(value):
        return None
    text = " ".join(str(value).strip().split())
    if not text:
        return None
    return text.title()


def _normalize_quarter(value):
    if pd.isna(value):
        return None
    text = "".join(str(value).strip().split()).upper()
    if not text:
        return None
    return text.replace("QUARTER", "Q")


def _load_catalog_rows():
    tables = pd.read_html(CATALOG_PDF)
    assert tables, "Could not parse product catalog PDF"
    catalog = pd.concat(tables, ignore_index=True)
    catalog.columns = [str(c).strip() for c in catalog.columns]
    return catalog


def build_expected_frame():
    catalog = _load_catalog_rows()
    sales = pd.read_excel(SALES_XLSX, dtype=object)

    catalog.columns = [str(c).strip() for c in catalog.columns]
    sales.columns = [str(c).strip() for c in sales.columns]

    product_col = _find_column(catalog.columns, "product", "id")
    name_col = _find_column(catalog.columns, "product", "name")
    category_col = _find_column(catalog.columns, "category")
    cost_col = _find_column(catalog.columns, "unit", "cost")
    catalog_price_col = _find_column(catalog.columns, "unit", "price")

    sales_product_col = _find_column(sales.columns, "product", "id")
    region_col = _find_column(sales.columns, "region")
    quantity_col = _find_column(sales.columns, "quantity")
    tx_price_col = _find_column(sales.columns, "unit", "price")
    month_col = _find_column(sales.columns, "month")
    quarter_col = _find_column(sales.columns, "quarter")

    catalog = catalog.rename(
        columns={
            product_col: "PRODUCT_ID",
            name_col: "PRODUCT_NAME",
            category_col: "CATEGORY",
            cost_col: "UNIT_COST",
            catalog_price_col: "CATALOG_UNIT_PRICE",
        }
    )
    sales = sales.rename(
        columns={
            sales_product_col: "PRODUCT_ID",
            region_col: "REGION",
            quantity_col: "QUANTITY",
            tx_price_col: "TRANSACTION_UNIT_PRICE",
            month_col: "MONTH",
            quarter_col: "QUARTER",
        }
    )

    sales["ORIGINAL_PRODUCT_ID"] = sales["PRODUCT_ID"]
    sales["PRODUCT_ID"] = sales["PRODUCT_ID"].map(lambda v: None if pd.isna(v) else str(v).strip())
    sales["REGION"] = sales["REGION"].map(_normalize_region)
    sales["MONTH"] = sales["MONTH"].map(lambda v: None if pd.isna(v) else str(v).strip().title())
    sales["QUARTER"] = sales["QUARTER"].map(_normalize_quarter)
    sales["TRANSACTION_UNIT_PRICE"] = pd.to_numeric(sales["TRANSACTION_UNIT_PRICE"], errors="coerce")
    sales["QUANTITY"] = pd.to_numeric(sales["QUANTITY"], errors="coerce")

    merged = sales.merge(catalog, on="PRODUCT_ID", how="left")

    cleaned_rows = []
    for row in merged.to_dict("records"):
        if row["PRODUCT_ID"] is None:
            continue
        if pd.isna(row["PRODUCT_NAME"]):
            continue
        qty = row["QUANTITY"]
        if pd.isna(qty) or qty <= 0:
            continue

        transaction_price = row["TRANSACTION_UNIT_PRICE"]
        catalog_price = pd.to_numeric(row["CATALOG_UNIT_PRICE"], errors="coerce")
        unit_cost = pd.to_numeric(row["UNIT_COST"], errors="coerce")

        if pd.notna(transaction_price):
            final_price = float(transaction_price)
            price_status = "transaction_override"
            action = "used_transaction_unit_price"
        else:
            final_price = float(catalog_price)
            price_status = "catalog_price"
            action = "filled_unit_price_from_catalog"

        cleaned_rows.append(
            {
                "PRODUCT_ID": row["PRODUCT_ID"],
                "PRODUCT_NAME": row["PRODUCT_NAME"],
                "CATEGORY": row["CATEGORY"],
                "REGION": row["REGION"],
                "MONTH": row["MONTH"],
                "QUARTER": row["QUARTER"],
                "QUANTITY": float(qty),
                "UNIT_PRICE": final_price,
                "UNIT_COST": float(unit_cost),
                "REVENUE": float(qty) * final_price,
                "PROFIT": float(qty) * (final_price - float(unit_cost)),
                "PRICE_STATUS": price_status,
                "CATALOG_MATCH_STATUS": "matched",
                "RECONCILIATION_ACTION": action,
            }
        )

    expected = pd.DataFrame(cleaned_rows)
    expected["MARGIN_PCT"] = expected["PROFIT"] / expected["REVENUE"]
    expected = expected.drop_duplicates().reset_index(drop=True)
    return expected


@pytest.fixture(scope="module")
def expected_frame():
    return build_expected_frame()


class TestSourceDataContent:
    def test_row_count_matches_expected(self, source_frame, expected_frame):
        assert len(source_frame) == len(expected_frame), f"Expected {len(expected_frame)} cleaned rows, got {len(source_frame)}"

    def test_category_values(self, source_frame):
        cat_col = _find_column(source_frame.columns, "category")
        cats = {str(v) for v in source_frame[cat_col].dropna().unique()}
        invalid = cats - VALID_CATEGORIES
        assert not invalid, f"Invalid categories: {invalid}"

    def test_region_values(self, source_frame):
        reg_col = _find_column(source_frame.columns, "region")
        regions = {str(v) for v in source_frame[reg_col].dropna().unique()}
        invalid = regions - VALID_REGIONS
        assert not invalid, f"Invalid regions: {invalid}"

    def test_reconciliation_metadata_values(self, source_frame):
        price_status_col = _find_column(source_frame.columns, "price", "status")
        match_status_col = _find_column(source_frame.columns, "catalog", "match")
        action_col = _find_column(source_frame.columns, "reconciliation", "action")
        assert set(source_frame[price_status_col].dropna().unique()).issubset(ALLOWED_PRICE_STATUS)
        assert set(source_frame[match_status_col].dropna().unique()).issubset(ALLOWED_MATCH_STATUS)
        assert set(source_frame[action_col].dropna().unique()).issubset(ALLOWED_ACTIONS)


class TestDataTransformations:
    def test_revenue_calculation(self, source_frame):
        qty_col = _find_column(source_frame.columns, "quantity")
        price_col = _find_column(source_frame.columns, "unit", "price")
        rev_col = _find_column(source_frame.columns, "revenue")
        for i, row in source_frame.head(80).iterrows():
            assert math.isclose(float(row[qty_col]) * float(row[price_col]), float(row[rev_col]), rel_tol=0, abs_tol=0.01), i

    def test_profit_calculation(self, source_frame):
        rev_col = _find_column(source_frame.columns, "revenue")
        qty_col = _find_column(source_frame.columns, "quantity")
        cost_col = _find_column(source_frame.columns, "unit", "cost")
        profit_col = _find_column(source_frame.columns, "profit")
        for i, row in source_frame.head(80).iterrows():
            expected = float(row[rev_col]) - float(row[qty_col]) * float(row[cost_col])
            assert math.isclose(expected, float(row[profit_col]), rel_tol=0, abs_tol=0.01), i

    def test_margin_range(self, source_frame):
        margin_col = _find_column(source_frame.columns, "margin")
        values = source_frame[margin_col].astype(float)
        assert ((values >= -1.0) & (values <= 1.0)).all()

    def test_source_data_matches_independent_expected(self, source_frame, expected_frame):
        ordered_columns = [
            "PRODUCT_ID",
            "PRODUCT_NAME",
            "CATEGORY",
            "REGION",
            "MONTH",
            "QUARTER",
            "QUANTITY",
            "UNIT_PRICE",
            "UNIT_COST",
            "REVENUE",
            "PROFIT",
            "MARGIN_PCT",
            "PRICE_STATUS",
            "CATALOG_MATCH_STATUS",
            "RECONCILIATION_ACTION",
        ]
        actual = source_frame[ordered_columns].copy()
        expected = expected_frame[ordered_columns].copy()
        for col in ["QUANTITY", "UNIT_PRICE", "UNIT_COST", "REVENUE", "PROFIT", "MARGIN_PCT"]:
            actual[col] = actual[col].astype(float).round(6)
            expected[col] = expected[col].astype(float).round(6)
        pd.testing.assert_frame_equal(actual.reset_index(drop=True), expected.reset_index(drop=True), check_dtype=False)

    def test_pivot_values_match_expected(self, source_frame, expected_frame):
        actual_revenue = source_frame.groupby(_find_column(source_frame.columns, "category"))[_find_column(source_frame.columns, "revenue")].sum().to_dict()
        expected_revenue = expected_frame.groupby("CATEGORY")["REVENUE"].sum().to_dict()
        assert actual_revenue == expected_revenue

        actual_units = source_frame.groupby(_find_column(source_frame.columns, "region"))[_find_column(source_frame.columns, "quantity")].sum().to_dict()
        expected_units = expected_frame.groupby("REGION")["QUANTITY"].sum().to_dict()
        assert actual_units == expected_units

        actual_counts = source_frame.groupby(_find_column(source_frame.columns, "category")).size().to_dict()
        expected_counts = expected_frame.groupby("CATEGORY").size().to_dict()
        assert actual_counts == expected_counts

        actual_matrix = defaultdict(dict)
        for (category, region), value in source_frame.groupby([
            _find_column(source_frame.columns, "category"),
            _find_column(source_frame.columns, "region"),
        ])[_find_column(source_frame.columns, "revenue")].sum().items():
            actual_matrix[category][region] = float(value)

        expected_matrix = defaultdict(dict)
        for (category, region), value in expected_frame.groupby(["CATEGORY", "REGION"])["REVENUE"].sum().items():
            expected_matrix[category][region] = float(value)

        assert dict(actual_matrix) == dict(expected_matrix)

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Revenue by Category"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
