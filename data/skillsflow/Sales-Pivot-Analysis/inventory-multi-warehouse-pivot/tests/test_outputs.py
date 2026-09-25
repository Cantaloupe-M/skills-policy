#!/usr/bin/env python3
"""Tests for Inventory Multi-Warehouse Pivot Table Analysis task."""
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

OUTPUT_FILE = "/root/inventory_report.xlsx"
PRODUCT_PDF = "/root/product_info.pdf"
WAREHOUSE_A_XLSX = "/root/warehouse_a_inventory.xlsx"
WAREHOUSE_B_XLSX = "/root/warehouse_b_inventory.xlsx"

PIVOT_SHEETS = [
    ("Stock by Category", "sum", None),
    ("Value by Warehouse", "sum", None),
    ("Items by Category", "count", None),
    ("Category Warehouse Matrix", "sum", "warehouse"),
]

REQUIRED_COLUMNS = [
    ("SKU", lambda h: "sku" in h),
    ("ITEM_NAME", lambda h: "item_name" in h or "itemname" in h.replace("_", "")),
    ("CATEGORY", lambda h: "category" in h),
    ("WEIGHT_KG", lambda h: "weight" in h and "kg" in h),
    ("REORDER_LEVEL", lambda h: "reorder" in h and "level" in h),
    ("WAREHOUSE", lambda h: "warehouse" in h),
    ("QUANTITY_ON_HAND", lambda h: ("quantity" in h and "hand" in h) or "qty" in h),
    ("UNIT_VALUE", lambda h: "unit_value" in h or "unitvalue" in h.replace("_", "")),
    ("TOTAL_VALUE", lambda h: "total_value" in h or "totalvalue" in h.replace("_", "")),
    ("TOTAL_WEIGHT", lambda h: "total_weight" in h or "totalweight" in h.replace("_", "")),
    ("REORDER_FLAG", lambda h: ("reorder" in h and "flag" in h) or h == "reorder_flag"),
    ("STOCK_STATUS", lambda h: "stock" in h and "status" in h),
    ("VALUE_TIER", lambda h: "value" in h and "tier" in h),
]

VALID_CATEGORIES = {"Raw Materials", "Components", "Finished Goods", "Packaging"}
VALID_WAREHOUSES = {"Warehouse-A", "Warehouse-B"}
VALID_REORDER_FLAG = {"Yes", "No"}
VALID_STOCK_STATUS = {"at_risk", "healthy"}
VALID_VALUE_TIERS = {"low", "medium", "high"}


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Category" in sheet_name:
            assert row_field and "category" in row_field.lower(), f"Row field should be CATEGORY, got '{row_field}'"
        elif "Warehouse" in sheet_name and "Category" not in sheet_name:
            assert row_field and "warehouse" in row_field.lower(), f"Row field should be WAREHOUSE, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and col_field in actual_col.lower(), f"Column field should be '{col_field}', got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_frame(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    raw_headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data_rows = [row for row in rows[1:] if any(cell is not None for cell in row)]
    return pd.DataFrame(data_rows, columns=raw_headers)


def _find_column(columns, *keywords):
    for col in columns:
        normalized = col.lower().replace("_", "").replace(" ", "")
        if all(keyword in normalized for keyword in keywords):
            return col
    raise AssertionError(f"Missing column with keywords {keywords}. Found: {list(columns)}")


def _normalize_sku(value):
    if pd.isna(value):
        return None
    return str(value).strip().upper().replace(" ", "")


def _value_tier(total_value):
    if total_value < 5000:
        return "low"
    if total_value < 20000:
        return "medium"
    return "high"


def build_expected_frame():
    product_tables = pd.read_html(PRODUCT_PDF)
    assert product_tables, "Could not parse product info PDF"
    products = pd.concat(product_tables, ignore_index=True)
    warehouse_a = pd.read_excel(WAREHOUSE_A_XLSX, dtype=object)
    warehouse_b = pd.read_excel(WAREHOUSE_B_XLSX, dtype=object)
    inventory = pd.concat([warehouse_a, warehouse_b], ignore_index=True)

    products.columns = [str(c).strip() for c in products.columns]
    inventory.columns = [str(c).strip() for c in inventory.columns]

    sku_col = _find_column(products.columns, "sku")
    item_col = _find_column(products.columns, "item", "name")
    category_col = _find_column(products.columns, "category")
    weight_col = _find_column(products.columns, "weight", "kg")
    reorder_col = _find_column(products.columns, "reorder", "level")

    inv_sku_col = _find_column(inventory.columns, "sku")
    warehouse_col = _find_column(inventory.columns, "warehouse")
    qty_col = _find_column(inventory.columns, "quantity", "hand")
    unit_value_col = _find_column(inventory.columns, "unit", "value")

    products = products.rename(
        columns={
            sku_col: "SKU",
            item_col: "ITEM_NAME",
            category_col: "CATEGORY",
            weight_col: "WEIGHT_KG",
            reorder_col: "REORDER_LEVEL",
        }
    )
    inventory = inventory.rename(
        columns={
            inv_sku_col: "SKU",
            warehouse_col: "WAREHOUSE",
            qty_col: "QUANTITY_ON_HAND",
            unit_value_col: "UNIT_VALUE",
        }
    )

    inventory["SKU"] = inventory["SKU"].map(_normalize_sku)
    inventory["QUANTITY_ON_HAND"] = pd.to_numeric(inventory["QUANTITY_ON_HAND"], errors="coerce")
    inventory["UNIT_VALUE"] = pd.to_numeric(inventory["UNIT_VALUE"], errors="coerce")
    products["SKU"] = products["SKU"].map(_normalize_sku)
    products["WEIGHT_KG"] = pd.to_numeric(products["WEIGHT_KG"], errors="coerce")
    products["REORDER_LEVEL"] = pd.to_numeric(products["REORDER_LEVEL"], errors="coerce")

    merged = inventory.merge(products, on="SKU", how="left")

    rows = []
    for row in merged.to_dict("records"):
        if row["SKU"] is None or pd.isna(row["ITEM_NAME"]):
            continue
        if pd.isna(row["QUANTITY_ON_HAND"]) or pd.isna(row["UNIT_VALUE"]):
            continue
        qty = float(row["QUANTITY_ON_HAND"])
        unit_value = float(row["UNIT_VALUE"])
        total_value = qty * unit_value
        rows.append(
            {
                "SKU": row["SKU"],
                "ITEM_NAME": row["ITEM_NAME"],
                "CATEGORY": row["CATEGORY"],
                "WEIGHT_KG": float(row["WEIGHT_KG"]),
                "REORDER_LEVEL": float(row["REORDER_LEVEL"]),
                "WAREHOUSE": row["WAREHOUSE"],
                "QUANTITY_ON_HAND": qty,
                "UNIT_VALUE": unit_value,
                "TOTAL_VALUE": total_value,
                "TOTAL_WEIGHT": qty * float(row["WEIGHT_KG"]),
                "REORDER_FLAG": "Yes" if qty < float(row["REORDER_LEVEL"]) else "No",
                "STOCK_STATUS": "at_risk" if qty < float(row["REORDER_LEVEL"]) else "healthy",
                "VALUE_TIER": _value_tier(total_value),
            }
        )

    expected = pd.DataFrame(rows)
    expected = expected.drop_duplicates().reset_index(drop=True)
    return expected


@pytest.fixture(scope="module")
def expected_frame():
    return build_expected_frame()


class TestSourceDataContent:
    def test_row_count(self, source_frame, expected_frame):
        assert len(source_frame) == len(expected_frame)

    def test_both_warehouses_present(self, source_frame):
        wh_col = _find_column(source_frame.columns, "warehouse")
        warehouses = {str(v) for v in source_frame[wh_col].dropna().unique()}
        assert warehouses == VALID_WAREHOUSES

    def test_category_values(self, source_frame):
        cat_col = _find_column(source_frame.columns, "category")
        cats = {str(v) for v in source_frame[cat_col].dropna().unique()}
        invalid = cats - VALID_CATEGORIES
        assert not invalid, f"Invalid categories: {invalid}"

    def test_status_columns(self, source_frame):
        reorder_flag_col = _find_column(source_frame.columns, "reorder", "flag")
        stock_status_col = _find_column(source_frame.columns, "stock", "status")
        value_tier_col = _find_column(source_frame.columns, "value", "tier")
        assert set(source_frame[reorder_flag_col].dropna().unique()).issubset(VALID_REORDER_FLAG)
        assert set(source_frame[stock_status_col].dropna().unique()).issubset(VALID_STOCK_STATUS)
        assert set(source_frame[value_tier_col].dropna().unique()).issubset(VALID_VALUE_TIERS)


class TestDataTransformations:
    def test_total_value_calculation(self, source_frame):
        qty_col = _find_column(source_frame.columns, "quantity", "hand")
        uval_col = _find_column(source_frame.columns, "unit", "value")
        tval_col = _find_column(source_frame.columns, "total", "value")
        for i, row in source_frame.head(120).iterrows():
            assert math.isclose(float(row[qty_col]) * float(row[uval_col]), float(row[tval_col]), rel_tol=0, abs_tol=0.01), i

    def test_reorder_flag_correctness(self, source_frame):
        qty_col = _find_column(source_frame.columns, "quantity", "hand")
        reorder_col = _find_column(source_frame.columns, "reorder", "level")
        flag_col = _find_column(source_frame.columns, "reorder", "flag")
        for i, row in source_frame.head(120).iterrows():
            expected = "Yes" if float(row[qty_col]) < float(row[reorder_col]) else "No"
            assert row[flag_col] == expected, i

    def test_source_matches_independent_expected(self, source_frame, expected_frame):
        ordered_columns = [
            "SKU",
            "ITEM_NAME",
            "CATEGORY",
            "WEIGHT_KG",
            "REORDER_LEVEL",
            "WAREHOUSE",
            "QUANTITY_ON_HAND",
            "UNIT_VALUE",
            "TOTAL_VALUE",
            "TOTAL_WEIGHT",
            "REORDER_FLAG",
            "STOCK_STATUS",
            "VALUE_TIER",
        ]
        actual = source_frame[ordered_columns].copy()
        expected = expected_frame[ordered_columns].copy()
        for col in ["WEIGHT_KG", "REORDER_LEVEL", "QUANTITY_ON_HAND", "UNIT_VALUE", "TOTAL_VALUE", "TOTAL_WEIGHT"]:
            actual[col] = actual[col].astype(float).round(6)
            expected[col] = expected[col].astype(float).round(6)
        pd.testing.assert_frame_equal(actual.reset_index(drop=True), expected.reset_index(drop=True), check_dtype=False)

    def test_pivot_aggregates_match_expected(self, source_frame, expected_frame):
        actual_stock = source_frame.groupby(_find_column(source_frame.columns, "category"))[_find_column(source_frame.columns, "quantity", "hand")].sum().round(6).to_dict()
        expected_stock = expected_frame.groupby("CATEGORY")["QUANTITY_ON_HAND"].sum().round(6).to_dict()
        assert actual_stock == expected_stock

        actual_value = source_frame.groupby(_find_column(source_frame.columns, "warehouse"))[_find_column(source_frame.columns, "total", "value")].sum().round(6).to_dict()
        expected_value = expected_frame.groupby("WAREHOUSE")["TOTAL_VALUE"].sum().round(6).to_dict()
        assert actual_value == expected_value

        actual_counts = source_frame.groupby(_find_column(source_frame.columns, "category")).size().to_dict()
        expected_counts = expected_frame.groupby("CATEGORY").size().to_dict()
        assert actual_counts == expected_counts

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Stock by Category"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
