#!/usr/bin/env python3
"""Tests for Student Performance Pivot Table Analysis task."""
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

OUTPUT_FILE = "/root/student_performance_report.xlsx"
ROSTER_PDF = "/root/student_roster.pdf"
GRADES_XLSX = "/root/course_grades.xlsx"

PIVOT_SHEETS = [
    ("Avg Score by Department", "average", None),
    ("Students by Department", "count", None),
    ("Credits by Semester", "sum", None),
    ("Department Semester Matrix", "average", "semester"),
]

REQUIRED_COLUMNS = [
    ("STUDENT_ID", lambda h: "student_id" in h or "studentid" in h.replace("_", "")),
    ("STUDENT_NAME", lambda h: "student_name" in h or "studentname" in h.replace("_", "")),
    ("DEPARTMENT", lambda h: "department" in h or "dept" in h),
    ("ENROLLMENT_YEAR", lambda h: "enrollment" in h and "year" in h),
    ("COURSE_NAME", lambda h: "course" in h),
    ("SEMESTER", lambda h: "semester" in h),
    ("SCORE", lambda h: "score" in h and "weighted" not in h),
    ("CREDITS", lambda h: "credit" in h),
    ("GRADE_BAND", lambda h: ("grade" in h and "band" in h) or h in {"grade_band", "gradeband"}),
    ("WEIGHTED_SCORE", lambda h: "weighted" in h),
    ("TERM_STATUS", lambda h: "term" in h and "status" in h),
    ("RETAKE_FLAG", lambda h: "retake" in h and "flag" in h),
]

VALID_DEPARTMENTS = {"Computer Science", "Mathematics", "Physics", "Biology", "Chemistry", "Economics", "English Literature"}
VALID_GRADE_BANDS = {"A", "B", "C", "D", "F"}
VALID_TERM_STATUS = {"standard", "summer_intensive"}
VALID_RETAKE_FLAG = {"Yes", "No"}


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Department" in sheet_name:
            assert row_field and "department" in row_field.lower(), f"Row field should be DEPARTMENT, got '{row_field}'"
        elif "Semester" in sheet_name and "Department" not in sheet_name:
            assert row_field and "semester" in row_field.lower(), f"Row field should be SEMESTER, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and col_field in actual_col.lower(), f"Column field should be '{col_field}', got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_frame(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    raw_headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data_rows = [row for row in rows[1:] if any(cell is not None for cell in row)]
    return pd.DataFrame(data_rows, columns=raw_headers)


def _find_column(columns, *keywords):
    for col in columns:
        normalized = col.lower().replace("_", "").replace(" ", "")
        if all(keyword in normalized for keyword in keywords):
            return col
    raise AssertionError(f"Missing column with keywords {keywords}. Found: {list(columns)}")


def _normalize_semester(value):
    if pd.isna(value):
        return None
    text = " ".join(str(value).strip().split())
    if not text:
        return None
    return text.title()


def _grade_band(score):
    if score >= 90:
        return "A"
    if score >= 80:
        return "B"
    if score >= 70:
        return "C"
    if score >= 60:
        return "D"
    return "F"


def build_expected_frame():
    roster_tables = pd.read_html(ROSTER_PDF)
    assert roster_tables, "Could not parse student roster PDF"
    roster = pd.concat(roster_tables, ignore_index=True)
    grades = pd.read_excel(GRADES_XLSX, dtype=object)

    roster.columns = [str(c).strip() for c in roster.columns]
    grades.columns = [str(c).strip() for c in grades.columns]

    roster_id_col = _find_column(roster.columns, "student", "id")
    name_col = _find_column(roster.columns, "student", "name")
    dept_col = _find_column(roster.columns, "department")
    enroll_col = _find_column(roster.columns, "enrollment", "year")

    grades_id_col = _find_column(grades.columns, "student", "id")
    course_col = _find_column(grades.columns, "course")
    semester_col = _find_column(grades.columns, "semester")
    score_col = _find_column(grades.columns, "score")
    credits_col = _find_column(grades.columns, "credit")

    roster = roster.rename(
        columns={
            roster_id_col: "STUDENT_ID",
            name_col: "STUDENT_NAME",
            dept_col: "DEPARTMENT",
            enroll_col: "ENROLLMENT_YEAR",
        }
    )
    grades = grades.rename(
        columns={
            grades_id_col: "STUDENT_ID",
            course_col: "COURSE_NAME",
            semester_col: "SEMESTER",
            score_col: "SCORE",
            credits_col: "CREDITS",
        }
    )

    grades["STUDENT_ID"] = grades["STUDENT_ID"].map(lambda v: None if pd.isna(v) else str(v).strip())
    grades["SEMESTER"] = grades["SEMESTER"].map(_normalize_semester)
    grades["SCORE"] = pd.to_numeric(grades["SCORE"], errors="coerce")
    grades["CREDITS"] = pd.to_numeric(grades["CREDITS"], errors="coerce")

    merged = grades.merge(roster, on="STUDENT_ID", how="left")

    rows = []
    for row in merged.to_dict("records"):
        if row["STUDENT_ID"] is None or pd.isna(row["STUDENT_NAME"]):
            continue
        if pd.isna(row["SCORE"]) or pd.isna(row["CREDITS"]):
            continue
        score = float(row["SCORE"])
        credits = float(row["CREDITS"])
        rows.append(
            {
                "STUDENT_ID": row["STUDENT_ID"],
                "STUDENT_NAME": row["STUDENT_NAME"],
                "DEPARTMENT": row["DEPARTMENT"],
                "ENROLLMENT_YEAR": int(float(row["ENROLLMENT_YEAR"])),
                "COURSE_NAME": row["COURSE_NAME"],
                "SEMESTER": row["SEMESTER"],
                "SCORE": score,
                "CREDITS": credits,
                "GRADE_BAND": _grade_band(score),
                "WEIGHTED_SCORE": score * credits,
                "TERM_STATUS": "summer_intensive" if "Summer" in str(row["SEMESTER"]) else "standard",
                "RETAKE_FLAG": "Yes" if score < 70 else "No",
            }
        )

    expected = pd.DataFrame(rows)
    expected = expected.drop_duplicates().reset_index(drop=True)
    return expected


@pytest.fixture(scope="module")
def expected_frame():
    return build_expected_frame()


class TestSourceDataContent:
    def test_row_count_matches_expected(self, source_frame, expected_frame):
        assert len(source_frame) == len(expected_frame)

    def test_department_values(self, source_frame):
        dept_col = _find_column(source_frame.columns, "department")
        depts = {str(v) for v in source_frame[dept_col].dropna().unique()}
        invalid = depts - VALID_DEPARTMENTS
        assert not invalid, f"Invalid departments: {invalid}"

    def test_grade_band_values(self, source_frame):
        band_col = _find_column(source_frame.columns, "grade", "band")
        bands = {str(v) for v in source_frame[band_col].dropna().unique()}
        invalid = bands - VALID_GRADE_BANDS
        assert not invalid, f"Invalid grade bands: {invalid}"

    def test_status_flags(self, source_frame):
        term_status_col = _find_column(source_frame.columns, "term", "status")
        retake_flag_col = _find_column(source_frame.columns, "retake", "flag")
        assert set(source_frame[term_status_col].dropna().unique()).issubset(VALID_TERM_STATUS)
        assert set(source_frame[retake_flag_col].dropna().unique()).issubset(VALID_RETAKE_FLAG)


class TestDataTransformations:
    def test_weighted_score_calculation(self, source_frame):
        score_col = _find_column(source_frame.columns, "score")
        credits_col = _find_column(source_frame.columns, "credit")
        weighted_col = _find_column(source_frame.columns, "weighted")
        for i, row in source_frame.head(120).iterrows():
            assert math.isclose(float(row[score_col]) * float(row[credits_col]), float(row[weighted_col]), rel_tol=0, abs_tol=0.01), i

    def test_grade_band_correctness(self, source_frame):
        score_col = _find_column(source_frame.columns, "score")
        band_col = _find_column(source_frame.columns, "grade", "band")
        for i, row in source_frame.head(120).iterrows():
            assert row[band_col] == _grade_band(float(row[score_col])), i

    def test_source_matches_independent_expected(self, source_frame, expected_frame):
        ordered_columns = [
            "STUDENT_ID",
            "STUDENT_NAME",
            "DEPARTMENT",
            "ENROLLMENT_YEAR",
            "COURSE_NAME",
            "SEMESTER",
            "SCORE",
            "CREDITS",
            "GRADE_BAND",
            "WEIGHTED_SCORE",
            "TERM_STATUS",
            "RETAKE_FLAG",
        ]
        actual = source_frame[ordered_columns].copy()
        expected = expected_frame[ordered_columns].copy()
        for col in ["ENROLLMENT_YEAR", "SCORE", "CREDITS", "WEIGHTED_SCORE"]:
            actual[col] = actual[col].astype(float).round(6)
            expected[col] = expected[col].astype(float).round(6)
        pd.testing.assert_frame_equal(actual.reset_index(drop=True), expected.reset_index(drop=True), check_dtype=False)

    def test_pivot_aggregates_match_expected(self, source_frame, expected_frame):
        actual_avg = source_frame.groupby(_find_column(source_frame.columns, "department"))[_find_column(source_frame.columns, "score")].mean().round(6).to_dict()
        expected_avg = expected_frame.groupby("DEPARTMENT")["SCORE"].mean().round(6).to_dict()
        assert actual_avg == expected_avg

        actual_counts = source_frame.groupby(_find_column(source_frame.columns, "department")).size().to_dict()
        expected_counts = expected_frame.groupby("DEPARTMENT").size().to_dict()
        assert actual_counts == expected_counts

        actual_credits = source_frame.groupby(_find_column(source_frame.columns, "semester"))[_find_column(source_frame.columns, "credit")].sum().round(6).to_dict()
        expected_credits = expected_frame.groupby("SEMESTER")["CREDITS"].sum().round(6).to_dict()
        assert actual_credits == expected_credits

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Avg Score by Department"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
