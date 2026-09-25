#!/usr/bin/env python3
"""Tests for Event Registration Pivot Table Analysis task."""
import pytest
from openpyxl import load_workbook

OUTPUT_FILE = "/root/registration_report.xlsx"

PIVOT_SHEETS = [
    ("Revenue by Track", "sum", None),
    ("Attendance by Venue", "count", None),
    ("Track RegType Matrix", "sum", "reg_type"),
    ("Events by Track", "count", None),
]

REQUIRED_COLUMNS = [
    ("REG_ID", lambda h: "reg_id" in h or "regid" in h.replace("_", "")),
    ("EVENT_ID", lambda h: "event_id" in h or "eventid" in h.replace("_", "")),
    ("EVENT_NAME", lambda h: "event_name" in h or "eventname" in h.replace("_", "")),
    ("TRACK", lambda h: "track" in h),
    ("VENUE", lambda h: "venue" in h),
    ("ATTENDEE_NAME", lambda h: "attendee" in h or "name" in h and "event" not in h),
    ("REG_TYPE", lambda h: "reg_type" in h or "regtype" in h.replace("_", "") or "registration_type" in h),
    ("AMOUNT_PAID", lambda h: "amount" in h or "paid" in h or "fee" in h),
    ("SOURCE", lambda h: "source" in h),
    ("IS_VIP", lambda h: "vip" in h),
    ("PRICE_TIER", lambda h: "price" in h and "tier" in h or h == "price_tier"),
]


@pytest.fixture(scope="module")
def workbook():
    return load_workbook(OUTPUT_FILE)


def _get_pivot_field_names(pivot):
    cache = pivot.cache
    if cache and cache.cacheFields:
        return [f.name for f in cache.cacheFields]
    return []


def _get_field_name_by_index(pivot, fields):
    field_names = _get_pivot_field_names(pivot)
    if fields and len(fields) > 0:
        idx = fields[0].x
        if idx is not None and 0 <= idx < len(field_names):
            return field_names[idx]
    return None


class TestPivotTableConfiguration:
    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_exists(self, workbook, sheet_name, expected_agg, col_field):
        assert sheet_name in workbook.sheetnames, f"Missing sheet '{sheet_name}'"
        pivots = workbook[sheet_name]._pivots
        assert len(pivots) > 0, f"No pivot table found in '{sheet_name}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_row_field(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        row_field = _get_field_name_by_index(pivot, pivot.rowFields)
        if "Track" in sheet_name:
            assert row_field and "track" in row_field.lower(), f"Row field should be TRACK, got '{row_field}'"
        elif "Venue" in sheet_name:
            assert row_field and "venue" in row_field.lower(), f"Row field should be VENUE, got '{row_field}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_aggregation(self, workbook, sheet_name, expected_agg, col_field):
        pivot = workbook[sheet_name]._pivots[0]
        data_field = pivot.dataFields[0]
        assert data_field.subtotal == expected_agg, f"Expected '{expected_agg}', got '{data_field.subtotal}'"

    @pytest.mark.parametrize("sheet_name,expected_agg,col_field", PIVOT_SHEETS)
    def test_pivot_col_field(self, workbook, sheet_name, expected_agg, col_field):
        if not col_field:
            pytest.skip(f"'{sheet_name}' is not a matrix pivot")
        pivot = workbook[sheet_name]._pivots[0]
        actual_col = _get_field_name_by_index(pivot, pivot.colFields)
        assert actual_col and ("reg" in actual_col.lower() or "type" in actual_col.lower()), \
            f"Column field should be REG_TYPE, got '{actual_col}'"


@pytest.fixture(scope="module")
def source_sheet(workbook):
    for name in workbook.sheetnames:
        if "source" in name.lower() or "data" in name.lower():
            return workbook[name]
    pytest.fail("No source data sheet found")


@pytest.fixture(scope="module")
def headers(source_sheet):
    first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(h).strip().lower() if h else "" for h in first_row]


class TestSourceDataSheet:
    @pytest.mark.parametrize("desc,match_fn", REQUIRED_COLUMNS)
    def test_has_required_column(self, headers, desc, match_fn):
        assert any(match_fn(h) for h in headers), f"Missing {desc} column. Found: {headers}"


@pytest.fixture(scope="module")
def source_data(source_sheet):
    rows = list(source_sheet.iter_rows(values_only=True))
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    data = [dict(zip(headers, row)) for row in rows[1:] if any(cell is not None for cell in row)]
    return data, headers


VALID_TRACKS = {"Data Science", "Web Development", "Cloud Infrastructure", "AI/ML", "Security"}
VALID_SOURCES = {"Online", "Walk-in"}
VALID_PRICE_TIERS = {"Free", "Budget", "Standard", "Premium"}


class TestSourceDataContent:
    def test_row_count(self, source_data):
        """Should have combined online + walk-in (minus orphans)."""
        data, _ = source_data
        # 900 online + ~400 walk-in minus ~20 orphans = ~1280
        assert 1000 <= len(data) <= 1500, f"Expected ~1280 rows, got {len(data)}"

    def test_both_sources_present(self, source_data):
        data, headers = source_data
        src_col = next((h for h in headers if "source" in h.lower()), None)
        sources = {row.get(src_col) for row in data if row.get(src_col)}
        assert "Online" in sources, "Missing Online source"
        assert "Walk-in" in sources, "Missing Walk-in source"

    def test_track_values(self, source_data):
        data, headers = source_data
        col = next((h for h in headers if "track" in h.lower()), None)
        vals = {row.get(col) for row in data if row.get(col)}
        invalid = vals - VALID_TRACKS
        assert not invalid, f"Invalid tracks: {invalid}"

    def test_orphan_events_excluded(self, source_data):
        """Event IDs not in catalog (9900+) should have been dropped."""
        data, headers = source_data
        eid_col = next((h for h in headers if "event_id" in h.lower() or "eventid" in h.lower()), None)
        if not eid_col:
            pytest.skip("No event_id column")
        eids = {int(float(row.get(eid_col))) for row in data if row.get(eid_col) is not None}
        orphans = {e for e in eids if e >= 9900}
        assert not orphans, f"Orphan event IDs should have been dropped: {orphans}"


class TestDataTransformations:
    def test_is_vip_values(self, source_data):
        data, headers = source_data
        vip_col = next((h for h in headers if "vip" in h.lower()), None)
        if not vip_col:
            pytest.skip("No IS_VIP column")
        vals = {row.get(vip_col) for row in data if row.get(vip_col) is not None}
        assert vals.issubset({"Yes", "No"}), f"Invalid IS_VIP values: {vals}"

    def test_is_vip_correctness(self, source_data):
        data, headers = source_data
        type_col = next((h for h in headers if "reg_type" in h.lower() or "regtype" in h.lower() or "registration_type" in h.lower()), None)
        vip_col = next((h for h in headers if "vip" in h.lower()), None)
        if not all([type_col, vip_col]):
            pytest.skip("Missing columns")
        errors = []
        for i, row in enumerate(data[:100]):
            rtype, vip = row.get(type_col), row.get(vip_col)
            if rtype and vip:
                expected = "Yes" if rtype == "VIP" else "No"
                if vip != expected:
                    errors.append(f"Row {i+2}: type={rtype}, got IS_VIP={vip}")
        assert not errors, f"IS_VIP errors:\n" + "\n".join(errors[:5])

    def test_price_tier_values(self, source_data):
        data, headers = source_data
        tier_col = next((h for h in headers if "price" in h.lower() and "tier" in h.lower()), None)
        if not tier_col:
            pytest.skip("No PRICE_TIER column")
        vals = {row.get(tier_col) for row in data if row.get(tier_col) is not None}
        invalid = vals - VALID_PRICE_TIERS
        assert not invalid, f"Invalid price tiers: {invalid}"

    def test_pivot_cache_has_fields(self, workbook):
        pivot = workbook["Revenue by Track"]._pivots[0]
        assert len(pivot.cache.cacheFields) > 0
