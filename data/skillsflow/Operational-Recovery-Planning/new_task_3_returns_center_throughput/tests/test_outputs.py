import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CONFIG = json.loads(r"""{
  "input_path": "/root/returns_intake_reference.xlsx",
  "input_sheet": "Returns",
  "input_week_label": "Week",
  "input_demand_labels": ["Standard Return Intake Hours", "Exception Review Hours"],
  "template_input": null,
  "output_workbook": "/root/returns_throughput_plan.xlsx",
  "output_summary": "/root/returns_throughput_summary.txt",
  "plan_sheet": "Plan",
  "plan_headers": ["Week", "Processing Days", "Forecast Return Intake (Work Hrs)", "Weekly Processing Capacity (Work Hrs)", "Start-of-Week Return Queue (Work Hrs)", "End-of-Week Return Queue/Buffer (Work Hrs)", "Flex Shift Hours"],
  "required_sheetnames": ["Plan"],
  "template_preserve_cells": [],
  "week_start": 3,
  "week_end": 45,
  "initial_total": 467.2,
  "rate_per_day": 32.0,
  "overtime_per_extra_day": 9.0,
  "demand_threshold_for_4": 128.0,
  "summary_template": "Project CleanSweep hits the 5-day Milestone in Week {fw5} and the 4-day Milestone in Week {fw4}. Keep flex shifts focused on backlog removal, then hold the lighter pattern once intake fits 4-day capacity.",
  "summary_required_phrases": ["Project CleanSweep", "Milestone"]
}""")


def normalize_label(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip().lower()


def load_reference_demands(cfg: dict[str, Any]) -> dict[int, float]:
    wb = load_workbook(cfg['input_path'], data_only=True)
    ws = wb[cfg['input_sheet']] if cfg['input_sheet'] in wb.sheetnames else wb.active
    week_row = None
    demand_rows: dict[str, tuple[Any, ...]] = {}
    wanted = {normalize_label(label): label for label in cfg['input_demand_labels']}
    week_label = normalize_label(cfg['input_week_label'])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        label = normalize_label(row[0] if row else None)
        if label == week_label:
            week_row = row
        elif label in wanted:
            demand_rows[label] = row
    assert week_row is not None, f"Missing week row '{cfg['input_week_label']}'"
    missing = [label for label in wanted if label not in demand_rows]
    assert not missing, f"Missing demand row(s): {missing}"
    week_to_demand: dict[int, float] = {}
    for idx, week_val in enumerate(week_row[1:], start=1):
        if week_val is None:
            continue
        week = int(round(float(week_val)))
        total = 0.0
        for label in wanted:
            row = demand_rows[label]
            total += float(row[idx])
        week_to_demand[week] = round(total, 2)
    expected = range(cfg['week_start'], cfg['week_end'] + 1)
    for w in expected:
        assert w in week_to_demand, f"Reference demand missing for week {w}"
    return week_to_demand


def load_plan_rows(cfg: dict[str, Any]) -> list[tuple[int, int, float, float, float, float, float]]:
    plan_file = Path(cfg['output_workbook'])
    assert plan_file.exists(), f"Output workbook missing: {plan_file}"
    wb = load_workbook(plan_file, data_only=True)
    assert set(wb.sheetnames) == set(cfg['required_sheetnames']), f"Sheet names mismatch: {wb.sheetnames}"
    ws = wb[cfg['plan_sheet']]
    headers = [ws.cell(1, i).value for i in range(1, 8)]
    assert set(headers) == set(cfg['plan_headers']), f"Header mismatch: {headers}"
    rows: list[tuple[int, int, float, float, float, float, float]] = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if all(v is None for v in vals):
            continue
        rows.append((
            int(round(float(vals[0]))),
            int(round(float(vals[1]))),
            float(vals[2]),
            float(vals[3]),
            float(vals[4]),
            float(vals[5]),
            float(vals[6]),
        ))
    expected_count = cfg['week_end'] - cfg['week_start'] + 1
    assert len(rows) == expected_count, f"Expected {expected_count} rows, got {len(rows)}"
    return rows


def expected_days(cfg: dict[str, Any], calc_start: float, demand: float) -> int:
    if max(0.0, calc_start) > 0.01:
        if calc_start + demand - (cfg['rate_per_day'] * 5) <= 0.0:
            return 5
        return 6
    if demand <= cfg['demand_threshold_for_4']:
        return 4
    return 5


def test_plan_math_and_policy():
    ref = load_reference_demands(CONFIG)
    rows = load_plan_rows(CONFIG)
    expected_weeks = list(range(CONFIG['week_start'], CONFIG['week_end'] + 1))
    actual_weeks = [r[0] for r in rows]
    assert actual_weeks == expected_weeks, 'Weeks must be continuous and ordered.'
    calc_start = round(CONFIG['initial_total'] - ref[CONFIG['week_start']], 2)
    prev_end = None
    seen = set()
    for week, days, demand, capacity, start_queue, end_queue, overtime in rows:
        seen.add(days)
        assert days in {4, 5, 6}, f"Week {week}: invalid days {days}"
        assert abs(demand - ref[week]) <= 0.01, f"Week {week}: demand mismatch"
        if week == CONFIG['week_start']:
            assert abs(start_queue + demand - CONFIG['initial_total']) <= 0.01
        if prev_end is not None:
            assert abs(calc_start - prev_end) <= 0.01
        assert abs(start_queue - max(0.0, calc_start)) <= 0.01
        assert abs(capacity - CONFIG['rate_per_day'] * days) <= 0.01
        expected_end = round(calc_start + demand - capacity, 2)
        assert abs(end_queue - expected_end) <= 0.01
        expected_ot = round(CONFIG['overtime_per_extra_day'] * max(0, days - 4), 2)
        assert abs(overtime - expected_ot) <= 0.01
        assert days == expected_days(CONFIG, calc_start, demand)
        prev_end = end_queue
        calc_start = end_queue
    assert {4, 5, 6}.issubset(seen), 'Plan should include at least one 4-day, 5-day, and 6-day week.'


def test_summary_file():
    summary_file = Path(CONFIG['output_summary'])
    assert summary_file.exists(), f"Summary file missing: {summary_file}"
    rows = load_plan_rows(CONFIG)
    fw5 = next((w for w, d, *_ in rows if d == 5), None)
    fw4 = next((w for w, d, *_ in rows if d == 4), None)
    w5 = str(fw5) if fw5 is not None else 'N/A'
    w4 = str(fw4) if fw4 is not None else 'N/A'

    content = summary_file.read_text(encoding='utf-8').strip()
    lines = content.splitlines()
    assert abs(len(lines) - 3) <= 1, 'Summary must contain approximately 3 lines (2-4).'

    m5 = re.fullmatch(r'First_Week_5_Days:\s*(\d+|N/A)', lines[0].strip())
    m4 = re.fullmatch(r'First_Week_4_Days:\s*(\d+|N/A)', lines[1].strip())
    assert m5, 'Line 1 must match First_Week_5_Days: <week-or-N/A>'
    assert m4, 'Line 2 must match First_Week_4_Days: <week-or-N/A>'
    assert m5.group(1) == w5
    assert m4.group(1) == w4

    assert lines[2].startswith('Summary: ')
    text = lines[2][len('Summary: '):].strip()
    words = re.findall(r'\S+', text)
    assert len(words) <= 60, f'Summary exceeds 60 words ({len(words)})'
    assert len(re.findall(r'[.!?]', text)) <= 3, 'Summary exceeds 3 sentences.'
    assert w5 in text, 'Summary must mention 5-day week.'
    assert w4 in text, 'Summary must mention 4-day week.'
    for phrase in CONFIG['summary_required_phrases']:
        assert phrase in text, f'Summary must include phrase: {phrase}'
