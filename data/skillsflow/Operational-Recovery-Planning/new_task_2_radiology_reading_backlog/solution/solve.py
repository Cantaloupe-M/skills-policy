from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

CONFIG = json.loads(r"""{
  "input_path": "/root/radiology_reading_reference.xlsx",
  "input_sheet": "Reads",
  "input_week_label": "Week",
  "input_demand_labels": ["Reading Load Forecast Total"],
  "template_input": null,
  "output_workbook": "/root/radiology_recovery_plan.xlsx",
  "output_summary": "/root/radiology_recovery_summary.txt",
  "plan_sheet": "Plan",
  "plan_headers": ["Week", "Radiologist Days", "Forecast Reading Load (Scan Hrs)", "Weekly Reading Capacity (Scan Hrs)", "Start-of-Week Reading Backlog (Scan Hrs)", "End-of-Week Reading Backlog/Buffer (Scan Hrs)", "Surge Premium Hours"],
  "required_sheetnames": ["Plan"],
  "template_preserve_cells": [],
  "week_start": 6,
  "week_end": 54,
  "initial_total": 468.55,
  "rate_per_day": 26.0,
  "overtime_per_extra_day": 6.0,
  "demand_threshold_for_4": 104.0,
  "summary_template": "Project PulseLift reaches the 5-day Milestone in Week {fw5} and the 4-day Milestone in Week {fw4}. Keep surge coverage tied to backlog clearance, then settle into the lighter cadence once demand fits 4-day capacity.",
  "summary_required_phrases": ["Project PulseLift", "Milestone"]
}""")


def normalize_label(value: object) -> str:
    if value is None:
        return ''
    return str(value).strip().lower()


def load_weekly_demands(cfg: dict[str, Any]) -> list[tuple[int, float]]:
    wb = load_workbook(cfg['input_path'], data_only=True)
    ws = wb[cfg['input_sheet']] if cfg['input_sheet'] in wb.sheetnames else wb.active
    week_row = None
    demand_rows: dict[str, tuple[Any, ...]] = {}
    wanted = {normalize_label(label): label for label in cfg['input_demand_labels']}
    week_label = normalize_label(cfg['input_week_label'])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        label = normalize_label(row[0] if row else None)
        if label == week_label:
            week_row = row
        elif label in wanted:
            demand_rows[label] = row
    if week_row is None:
        raise ValueError(f"Missing week row '{cfg['input_week_label']}'.")
    missing = [label for label in wanted if label not in demand_rows]
    if missing:
        raise ValueError(f"Missing demand row(s): {missing}")
    week_to_demand: dict[int, float] = {}
    for idx, week_val in enumerate(week_row[1:], start=1):
        if week_val is None:
            continue
        week = int(round(float(week_val)))
        total = 0.0
        for label in wanted:
            row = demand_rows[label]
            total += float(row[idx])
        week_to_demand[week] = round(total, 2)
    expected = range(cfg['week_start'], cfg['week_end'] + 1)
    missing_weeks = [w for w in expected if w not in week_to_demand]
    if missing_weeks:
        raise ValueError(f"Missing weekly demand values for: {missing_weeks}")
    return [(w, week_to_demand[w]) for w in expected]


def choose_days(cfg: dict[str, Any], calc_start: float, demand: float) -> int:
    start_queue = max(0.0, calc_start)
    if start_queue > 0.01:
        for days in (5, 6):
            if calc_start + demand - (cfg['rate_per_day'] * days) <= 0.0:
                return days
        return 6
    if demand <= cfg['demand_threshold_for_4']:
        return 4
    return 5


def build_rows(cfg: dict[str, Any]) -> list[tuple[int, int, float, float, float, float, float]]:
    demands = load_weekly_demands(cfg)
    calc_start = round(cfg['initial_total'] - demands[0][1], 2)
    rows: list[tuple[int, int, float, float, float, float, float]] = []
    for week, demand in demands:
        days = choose_days(cfg, calc_start, demand)
        capacity = round(cfg['rate_per_day'] * days, 2)
        start_queue = round(max(0.0, calc_start), 2)
        end_queue = round(calc_start + demand - capacity, 2)
        overtime = round(cfg['overtime_per_extra_day'] * max(0, days - 4), 2)
        rows.append((week, days, round(demand, 2), capacity, start_queue, end_queue, overtime))
        calc_start = end_queue
    return rows


def build_workbook(cfg: dict[str, Any], rows: list[tuple[int, int, float, float, float, float, float]]) -> None:
    out_path = Path(cfg['output_workbook'])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    template = cfg.get('template_input')
    if template:
        wb = load_workbook(template)
    else:
        wb = Workbook()
        wb.active.title = cfg['plan_sheet']
    for name in cfg['required_sheetnames']:
        if name not in wb.sheetnames:
            wb.create_sheet(name)
    wb._sheets = [wb[name] for name in cfg['required_sheetnames']]
    plan = wb[cfg['plan_sheet']]
    plan.delete_rows(1, plan.max_row)
    plan.append(cfg['plan_headers'])
    for r in rows:
        plan.append(list(r))
    wb.save(out_path)


def build_summary(cfg: dict[str, Any], rows: list[tuple[int, int, float, float, float, float, float]]) -> None:
    fw5 = next((w for w, d, *_ in rows if d == 5), None)
    fw4 = next((w for w, d, *_ in rows if d == 4), None)
    w5 = str(fw5) if fw5 is not None else 'N/A'
    w4 = str(fw4) if fw4 is not None else 'N/A'
    summary = cfg['summary_template'].format(fw5=w5, fw4=w4)
    out_path = Path(cfg['output_summary'])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(
        '\n'.join([
            f'First_Week_5_Days: {w5}',
            f'First_Week_4_Days: {w4}',
            f'Summary: {summary}',
        ]),
        encoding='utf-8',
    )


def run(cfg_override: dict[str, Any] | None = None) -> list[tuple[int, int, float, float, float, float, float]]:
    cfg = dict(CONFIG)
    if cfg_override:
        cfg.update(cfg_override)
    rows = build_rows(cfg)
    build_workbook(cfg, rows)
    build_summary(cfg, rows)
    return rows


def main() -> None:
    run()


if __name__ == '__main__':
    main()
