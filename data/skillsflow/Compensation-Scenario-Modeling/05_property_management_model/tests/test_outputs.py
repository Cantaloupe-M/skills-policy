import pytest
import openpyxl
import os

OUTPUT_FILE = '/root/Property_Management.xlsx'


class TestSheetStructure:
    def test_file_exists(self):
        assert os.path.exists(OUTPUT_FILE), f"Output file not found: {OUTPUT_FILE}"

    def test_sheet_names(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        expected = ['Summary', 'Assumptions', 'Building Specs', 'Roster', 'Calculations --->',
                    'EE Calcs (Current)', 'EE Calcs (Yr+1)', 'EE Calcs (Yr+2)']
        assert wb.sheetnames == expected, f"Expected {expected}, got {wb.sheetnames}"

    def test_named_ranges_count(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        assert len(wb.defined_names) >= 46, \
            f"Expected >=46 named ranges, got {len(wb.defined_names)}"

    def test_summary_title(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb['Summary']
        assert ws['B1'].value is not None, "Summary B1 should contain organization name"
        assert ws['C5'].value == 52000, f"Expected C5=52000, got {ws['C5'].value}"
        assert ws['D5'].value == '=C5*1.03', f"Expected D5='=C5*1.03'"

    def test_ee_calcs_row_count(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb['EE Calcs (Current)']
        # 87 data rows + 1 totals = 88 rows starting at row 4
        data_rows = sum(1 for r in ws.iter_rows(min_row=4, max_row=100, min_col=2, max_col=2) if r[0].value is not None)
        assert data_rows >= 87, f"Expected >=87 data rows, got {data_rows}"

    def test_quarterly_totals_row_91(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb['EE Calcs (Current)']
        h91 = ws.cell(row=91, column=8).value
        assert h91 is not None, "Row 91 (quarterly totals) should exist in EE Calcs"

    def test_summary_references_row_91(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb['Summary']
        # Check that Summary row 33 references row 91 of EE Calcs
        for col in [3, 4, 5]:
            v = ws.cell(row=33, column=col).value
            assert v is not None and '91' in str(v), \
                f"Summary row 33 should reference EE Calcs row 91, got {v}"

    def test_roster_count(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws = wb['Roster']
        count = sum(1 for r in ws.iter_rows(min_row=5, max_row=200, min_col=2, max_col=2) if r[0].value is not None)
        assert count >= 87, f"Expected >=87 roster entries, got {count}"

    def test_yoy_service_progression(self):
        wb = openpyxl.load_workbook(OUTPUT_FILE)
        ws1 = wb['EE Calcs (Current)']
        ws2 = wb['EE Calcs (Yr+1)']
        for r in range(4, 91):
            v1 = ws1.cell(row=r, column=6).value
            v2 = ws2.cell(row=r, column=6).value
            if v1 is not None and v2 is not None:
                assert v2 == v1 + 1, f"Yr+1 service should be +1"
                break
