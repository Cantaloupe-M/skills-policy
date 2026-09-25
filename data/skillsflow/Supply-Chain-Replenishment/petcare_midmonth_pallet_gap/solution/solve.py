#!/usr/bin/env python3
"""Solver for petcare_midmonth_pallet_gap."""
import math
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook

EPS = 1e-9


def normalize_text(value: Any) -> str:
    return str(value or '').strip().upper()


def to_number(value: Any) -> float:
    if value in (None, ''):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(',', '').strip()
    return float(text) if text else 0.0


def parse_date(value: Any) -> date | None:
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def to_iso(value: date | None) -> str:
    return value.isoformat() if value else ''


def diff_days(start: date, end: date) -> int:
    return (end - start).days


def add_days(start: date, days: int) -> date:
    return start + timedelta(days=days)


def round_to(value: float, decimals: int = 4) -> float:
    factor = 10 ** decimals
    return round(value * factor) / factor


def main():
    input_root = Path(sys.argv[1])
    output_root = Path(sys.argv[2])

    wb = load_workbook(input_root / 'Petcare_Midmonth_Snapshot.xlsx', data_only=True)
    stock_ws = wb['Current Stock']
    arr_ws = wb['Expected Arrivals']
    cfg_ws = wb['Pallet Guide']

    as_of = parse_date(stock_ws['B1'].value)
    horizon = parse_date(stock_ws['D1'].value)
    planning_days = diff_days(as_of, horizon)
    cases_per_pallet = to_number(cfg_ws['A2'].value)

    inventory = []
    for r in range(4, stock_ws.max_row + 1):
        sku = normalize_text(stock_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        units = to_number(stock_ws.cell(row=r, column=2).value)
        rate = to_number(stock_ws.cell(row=r, column=3).value)
        inventory.append((sku, units, rate))

    valid_statuses = {'COMMITTED', 'ARRANGED'}
    inbound_by_sku: Dict[str, List[Tuple[date, float]]] = {}
    for r in range(2, arr_ws.max_row + 1):
        sku = normalize_text(arr_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        d = parse_date(arr_ws.cell(row=r, column=2).value)
        if d is None:
            continue
        cases = to_number(arr_ws.cell(row=r, column=3).value)
        status = normalize_text(arr_ws.cell(row=r, column=4).value)
        if status not in valid_statuses:
            continue
        if sku not in inbound_by_sku:
            inbound_by_sku[sku] = []
        inbound_by_sku[sku].append((d, cases))

    detail_rows = []
    action_rows = []

    for sku, units, rate in inventory:
        inbound = inbound_by_sku.get(sku, [])
        inbound_in_horizon = sum(c for d, c in inbound if d <= horizon)
        earliest = min((d for d, c in inbound if d <= horizon), default=None)
        current_doh = units / rate if rate > 0 else None
        proj_oos = add_days(as_of, int(current_doh)) if current_doh is not None else None
        delivered_doh = (units + inbound_in_horizon) / rate if rate > 0 else None
        remaining_demand = rate * planning_days
        additional = max(0.0, remaining_demand - units - inbound_in_horizon) if rate > 0 else 0.0
        pallets = math.ceil(additional / cases_per_pallet) if additional > 0 else 0
        req_date = proj_oos if pallets > 0 else None
        earlier = pallets > 0 and (earliest is None or (req_date and req_date < earliest))

        detail_rows.append([
            sku,
            units,
            rate,
            round_to(current_doh) if current_doh is not None else '',
            to_iso(proj_oos),
            inbound_in_horizon,
            round_to(delivered_doh) if delivered_doh is not None else '',
            round_to(remaining_demand),
            round_to(additional),
            pallets,
            to_iso(req_date),
            earlier,
        ])
        if pallets > 0:
            action_rows.append([sku, to_iso(req_date), pallets, round_to(additional), earlier])

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = 'SKU_Coverage'

    out_ws['A1'] = 'Field'
    out_ws['B1'] = 'Value'
    out_ws['A2'] = 'AsOfDate'
    out_ws['B2'] = to_iso(as_of)
    out_ws['A3'] = 'HorizonEnd'
    out_ws['B3'] = to_iso(horizon)
    out_ws['A4'] = 'PlanningDays'
    out_ws['B4'] = planning_days

    headers = [
        'SKU', 'Units_On_Hand', 'Daily_Rate_Units_Per_Day', 'Current_Days_On_Hand',
        'Projected_OOS_Date', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand',
        'Remaining_Demand_Units', 'Additional_Units_Needed', 'Pallets_Required',
        'Required_Delivery_Date', 'Earlier_Delivery_Required',
    ]
    for i, h in enumerate(headers, 1):
        out_ws.cell(row=6, column=i, value=h)

    for ridx, row in enumerate(detail_rows, 7):
        for cidx, val in enumerate(row, 1):
            out_ws.cell(row=ridx, column=cidx, value=val)

    act_ws = out_wb.create_sheet('Pallet_Gap_List')
    act_headers = ['SKU', 'Required_Delivery_Date', 'Pallets_Required', 'Additional_Units_Needed', 'Earlier_Delivery_Required']
    for i, h in enumerate(act_headers, 1):
        act_ws.cell(row=1, column=i, value=h)
    for ridx, row in enumerate(action_rows, 2):
        for cidx, val in enumerate(row, 1):
            act_ws.cell(row=ridx, column=cidx, value=val)

    output_file = output_root / 'petcare_midmonth_pallet_gap.xlsx'
    out_wb.save(output_file)

    print('Wrote', output_file)


if __name__ == '__main__':
    main()
