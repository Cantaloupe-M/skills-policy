#!/usr/bin/env python3
"""Solver for clinic_branch_transfer_gap."""
import math
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook

EPS = 1e-9


def normalize_text(value: Any) -> str:
    return str(value or '').strip().upper()


def to_number(value: Any) -> float:
    if value in (None, ''):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(',', '').strip()
    return float(text) if text else 0.0


def parse_date(value: Any) -> date | None:
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def to_iso(value: date | None) -> str:
    return value.isoformat() if value else ''


def diff_days(start: date, end: date) -> int:
    return (end - start).days


def add_days(start: date, days: int) -> date:
    return start + timedelta(days=days)


def round_to(value: float, decimals: int = 4) -> float:
    factor = 10 ** decimals
    return round(value * factor) / factor


def main():
    input_root = Path(sys.argv[1])
    output_root = Path(sys.argv[2])

    wbi = load_workbook(input_root / 'Clinic_Branch_Inventory.xlsx', data_only=True)
    stock_ws = wbi['Branch Stock']
    wbt = load_workbook(input_root / 'Clinic_Transfer_Schedule.xlsx', data_only=True)
    tr_ws = wbt['Planned Transfers']

    as_of = parse_date(stock_ws['B1'].value)
    horizon = parse_date(stock_ws['D1'].value)
    planning_days = diff_days(as_of, horizon)
    cases_per_pallet = 48

    inventory = []
    for r in range(4, stock_ws.max_row + 1):
        branch = normalize_text(stock_ws.cell(row=r, column=1).value)
        item = normalize_text(stock_ws.cell(row=r, column=2).value)
        if not branch or not item:
            continue
        units = to_number(stock_ws.cell(row=r, column=3).value)
        rate = to_number(stock_ws.cell(row=r, column=4).value)
        inventory.append((branch, item, units, rate))

    raw_transfers = []
    for r in range(2, tr_ws.max_row + 1):
        tid = normalize_text(tr_ws.cell(row=r, column=1).value)
        if not tid:
            continue
        branch = normalize_text(tr_ws.cell(row=r, column=2).value)
        item = normalize_text(tr_ws.cell(row=r, column=3).value)
        d = parse_date(tr_ws.cell(row=r, column=4).value)
        units = to_number(tr_ws.cell(row=r, column=5).value)
        status = normalize_text(tr_ws.cell(row=r, column=6).value)
        raw_transfers.append((tid, branch, item, d, units, status))

    by_tid: Dict[str, Tuple[str, str, date | None, float, str]] = {}
    for tid, branch, item, d, units, status in raw_transfers:
        if d is None:
            continue
        prev = by_tid.get(tid)
        if prev is None or d > prev[2]:
            by_tid[tid] = (branch, item, d, units, status)

    confirmed_statuses = {'CONFIRMED'}
    inbound_by_key: Dict[Tuple[str, str], List[Tuple[date, float]]] = {}
    for tid, (branch, item, d, units, status) in by_tid.items():
        if status not in confirmed_statuses:
            continue
        key = (branch, item)
        if key not in inbound_by_key:
            inbound_by_key[key] = []
        inbound_by_key[key].append((d, units))

    detail_rows = []
    action_rows = []

    for branch, item, units, rate in inventory:
        key = (branch, item)
        inbound = inbound_by_key.get(key, [])
        inbound_in_horizon = sum(u for d, u in inbound if d <= horizon)
        earliest = min((d for d, u in inbound if d <= horizon), default=None)
        current_doh = units / rate if rate > 0 else None
        proj_oos = add_days(as_of, int(current_doh)) if current_doh is not None else None
        delivered_doh = (units + inbound_in_horizon) / rate if rate > 0 else None
        remaining_demand = rate * planning_days
        additional = max(0.0, remaining_demand - units - inbound_in_horizon) if rate > 0 else 0.0
        pallets = math.ceil(additional / cases_per_pallet) if additional > 0 else 0
        req_date = proj_oos if pallets > 0 else None
        earlier = pallets > 0 and (earliest is None or (req_date and req_date < earliest))

        detail_rows.append([
            branch,
            item,
            units,
            rate,
            round_to(current_doh) if current_doh is not None else '',
            to_iso(proj_oos),
            inbound_in_horizon,
            round_to(delivered_doh) if delivered_doh is not None else '',
            round_to(remaining_demand),
            round_to(additional),
            pallets,
            to_iso(req_date),
            earlier,
        ])
        if pallets > 0:
            action_rows.append([branch, item, to_iso(req_date), pallets, round_to(additional), earlier])

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = 'Branch_Item_Coverage'

    out_ws['A1'] = 'Field'
    out_ws['B1'] = 'Value'
    out_ws['A2'] = 'AsOfDate'
    out_ws['B2'] = to_iso(as_of)
    out_ws['A3'] = 'HorizonEnd'
    out_ws['B3'] = to_iso(horizon)
    out_ws['A4'] = 'PlanningDays'
    out_ws['B4'] = planning_days

    headers = [
        'Branch', 'Item', 'Units_On_Hand', 'Daily_Use_Units_Per_Day', 'Current_Days_On_Hand',
        'Projected_OOS_Date', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand',
        'Remaining_Demand_Units', 'Additional_Units_Needed', 'Pallets_Required',
        'Required_Delivery_Date', 'Earlier_Delivery_Required',
    ]
    for i, h in enumerate(headers, 1):
        out_ws.cell(row=6, column=i, value=h)

    for ridx, row in enumerate(detail_rows, 7):
        for cidx, val in enumerate(row, 1):
            out_ws.cell(row=ridx, column=cidx, value=val)

    act_ws = out_wb.create_sheet('Transfer_Gap_List')
    act_headers = ['Branch', 'Item', 'Required_Delivery_Date', 'Pallets_Required', 'Additional_Units_Needed', 'Earlier_Delivery_Required']
    for i, h in enumerate(act_headers, 1):
        act_ws.cell(row=1, column=i, value=h)
    for ridx, row in enumerate(action_rows, 2):
        for cidx, val in enumerate(row, 1):
            act_ws.cell(row=ridx, column=cidx, value=val)

    output_file = output_root / 'clinic_branch_transfer_gap.xlsx'
    out_wb.save(output_file)

    print('Wrote', output_file)


if __name__ == '__main__':
    main()
