#!/usr/bin/env python3
"""Test outputs for frozen."""
import math
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook

EPS = 1e-9


def normalize_text(value: Any) -> str:
    return str(value or '').strip().upper()


def to_number(value: Any) -> float:
    if value in (None, ''):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(',', '').strip()
    return float(text) if text else 0.0


def parse_date(value: Any) -> date | None:
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def to_iso(value: date | None) -> str:
    return value.isoformat() if value else ''


def diff_days(start: date, end: date) -> int:
    return (end - start).days


def add_days(start: date, days: int) -> date:
    return start + timedelta(days=days)


def to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    s = normalize_text(value)
    return s in ('TRUE', 'T', 'YES', 'Y', '1')


def round_to(value: float, decimals: int = 4) -> float:
    factor = 10 ** decimals
    return round(value * factor) / factor


def assert_approx(actual: Any, expected: float, tolerance: float, msg: str):
    a = to_number(actual)
    assert abs(a - expected) <= tolerance, f'{msg}: actual={a}, expected={expected}'


def compute_expected_frozen(input_root: Path):
    """Compute expected outputs."""

    wbt = load_workbook(input_root / 'Frozen_Meal_Template.xlsx', data_only=True)
    guide_ws = wbt['Pallet Guide']
    wbs = load_workbook(input_root / 'Frozen_Meal_Current_Stock.xlsx', data_only=True)
    stock_ws = wbs['Stock']
    wbl = load_workbook(input_root / 'Frozen_Meal_Recovery_Log.xlsx', data_only=True)
    log_ws = wbl['Recovery Log']

    as_of = parse_date(stock_ws['B1'].value)
    horizon = parse_date(stock_ws['D1'].value)
    planning_days = diff_days(as_of, horizon)

    # Per-SKU pallet sizes
    pallet_by_sku = {}
    for r in range(2, guide_ws.max_row + 1):
        sku = normalize_text(guide_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        cpp = to_number(guide_ws.cell(row=r, column=2).value)
        pallet_by_sku[sku] = cpp

    inventory = []
    for r in range(4, stock_ws.max_row + 1):
        sku = normalize_text(stock_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        units = to_number(stock_ws.cell(row=r, column=2).value)
        rate = to_number(stock_ws.cell(row=r, column=3).value)
        inventory.append((sku, units, rate))

    # Dedupe Recovery Log by Load ID: keep max Revision
    raw_logs = []
    for r in range(2, log_ws.max_row + 1):
        lid = normalize_text(log_ws.cell(row=r, column=1).value)
        if not lid:
            continue
        rev = to_number(log_ws.cell(row=r, column=2).value)
        sku = normalize_text(log_ws.cell(row=r, column=3).value)
        d = parse_date(log_ws.cell(row=r, column=4).value)
        units = to_number(log_ws.cell(row=r, column=5).value)
        stage = normalize_text(log_ws.cell(row=r, column=6).value)
        raw_logs.append((lid, rev, sku, d, units, stage))

    by_lid = {}
    for lid, rev, sku, d, units, stage in raw_logs:
        if d is None:
            continue
        prev = by_lid.get(lid)
        if prev is None or rev > prev[0]:
            by_lid[lid] = (rev, sku, d, units, stage)

    valid_stages = {'BOOKED', 'LOADED'}
    inbound_by_sku = {}
    for lid, (rev, sku, d, units, stage) in by_lid.items():
        if stage not in valid_stages:
            continue
        if sku not in inbound_by_sku:
            inbound_by_sku[sku] = []
        inbound_by_sku[sku].append((d, units))

    detail_map = {}
    detail_order = []
    action_map = {}
    action_keys = set()

    for sku, units, rate in inventory:
        key = (sku,)
        detail_order.append(key)
        inbound = inbound_by_sku.get(sku, [])
        inbound_in_horizon = sum(u for d, u in inbound if d <= horizon)
        earliest = min((d for d, u in inbound if d <= horizon), default=None)
        current_doh = units / rate if rate > 0 else None
        proj_oos = add_days(as_of, int(current_doh)) if current_doh is not None else None
        delivered_doh = (units + inbound_in_horizon) / rate if rate > 0 else None
        remaining_demand = rate * planning_days
        additional = max(0.0, remaining_demand - units - inbound_in_horizon) if rate > 0 else 0.0
        cpp = pallet_by_sku.get(sku, 48)
        loads = math.ceil(additional / cpp) if additional > 0 else 0
        req_date = proj_oos if loads > 0 else None
        earlier = loads > 0 and (earliest is None or (req_date and req_date < earliest))

        row = {
            'SKU': sku,
            'Units_On_Hand': units,
            'Daily_Rate_Units_Per_Day': rate,
            'Current_Days_On_Hand': round_to(current_doh) if current_doh is not None else '',
            'Projected_OOS_Date': to_iso(proj_oos),
            'Inbound_Units_By_Horizon': inbound_in_horizon,
            'Delivered_Days_On_Hand': round_to(delivered_doh) if delivered_doh is not None else '',
            'Remaining_Demand_Units': round_to(remaining_demand),
            'Additional_Units_Needed': round_to(additional),
            'Loads_Required': loads,
            'Required_Delivery_Date': to_iso(req_date),
            'Earlier_Delivery_Required': earlier,
        }
        detail_map[key] = row
        if loads > 0:
            action_map[key] = {
                'SKU': sku,
                'Required_Delivery_Date': to_iso(req_date),
                'Loads_Required': loads,
                'Additional_Units_Needed': round_to(additional),
                'Earlier_Delivery_Required': earlier,
            }
            action_keys.add(key)

    return {
        'as_of': as_of,
        'horizon': horizon,
        'planning_days': planning_days,
        'detail_map': detail_map,
        'detail_order': detail_order,
        'action_map': action_map,
        'action_keys': action_keys,
    }



def read_output_frozen(output_file: Path):
    wb = load_workbook(output_file, data_only=True)
    assert set(wb.sheetnames) == set(['Instructions', 'Coverage_Detail', 'Recovery_Loads', 'Pallet Guide']), f'Sheet names mismatch: {wb.sheetnames}'
    detail_ws = wb['Coverage_Detail']
    action_ws = wb['Recovery_Loads']

    # Metadata
    assert detail_ws['A1'].value.strip().lower() == 'field'
    assert detail_ws['B1'].value.strip().lower() == 'value'
    as_of_raw = detail_ws['B2'].value
    as_of = parse_date(as_of_raw)
    assert as_of is not None, f'Invalid AsOfDate: {as_of_raw}'
    horizon_raw = detail_ws['B3'].value
    horizon = parse_date(horizon_raw)
    assert horizon is not None, f'Invalid HorizonEnd: {horizon_raw}'
    days_raw = detail_ws['B4'].value
    planning_days = int(to_number(days_raw))

    # Header row 6
    header = [detail_ws.cell(row=6, column=c).value for c in range(1, 12 + 1)]
    assert set(header) == set(['SKU', 'Units_On_Hand', 'Daily_Rate_Units_Per_Day', 'Current_Days_On_Hand', 'Projected_OOS_Date', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand', 'Remaining_Demand_Units', 'Additional_Units_Needed', 'Loads_Required', 'Required_Delivery_Date', 'Earlier_Delivery_Required']), f'Detail header mismatch: {header}'

    # Data rows
    detail_map = {}
    detail_order = []
    for r in range(7, detail_ws.max_row + 1):
        row = [detail_ws.cell(row=r, column=c).value for c in range(1, 12 + 1)]
        key = tuple(normalize_text(v) for v in row[:1])
        if not all(key):
            continue
        assert key not in detail_map, f'Duplicate key: {key}'
        obj = {header[i]: row[i] for i in range(len(header))}
        detail_map[key] = obj
        detail_order.append(key)

    # Action sheet
    action_header = [action_ws.cell(row=1, column=c).value for c in range(1, 5 + 1)]
    assert set(action_header) == set(['SKU', 'Required_Delivery_Date', 'Loads_Required', 'Additional_Units_Needed', 'Earlier_Delivery_Required']), f'Action header mismatch: {action_header}'
    action_map = {}
    for r in range(2, action_ws.max_row + 1):
        row = [action_ws.cell(row=r, column=c).value for c in range(1, 5 + 1)]
        key = tuple(normalize_text(v) for v in row[:1])
        if not all(key):
            continue
        obj = {action_header[i]: row[i] for i in range(len(action_header))}
        action_map[key] = obj

    return {
        'as_of': as_of,
        'horizon': horizon,
        'planning_days': planning_days,
        'detail_map': detail_map,
        'detail_order': detail_order,
        'action_map': action_map,
    }


def main():
    task_dir = Path(__file__).resolve().parent.parent

    in_harness = Path('/tests').exists() and Path('/root').exists()
    input_root = Path('/root') if in_harness else (task_dir / 'environment')
    output_file = Path('/root/frozen_meal_recovery_tracker.xlsx') if in_harness else (task_dir / 'frozen_meal_recovery_tracker.xlsx')

    assert output_file.exists(), f'Output file not found: {output_file}'

    expected = compute_expected_frozen(input_root)
    got = read_output_frozen(output_file)

    # Metadata checks
    assert got['as_of'] == expected['as_of'], f'AsOfDate mismatch'
    assert got['horizon'] == expected['horizon'], f'HorizonEnd mismatch'
    assert got['planning_days'] == expected['planning_days'], f'PlanningDays mismatch'

    # Row count and order
    assert len(got['detail_order']) == len(expected['detail_order']), 'Detail row count mismatch'
    assert set(got['detail_order']) == set(expected['detail_order']), 'Detail order mismatch'

    # Per-row checks
    for key in expected['detail_order']:
        exp_row = expected['detail_map'][key]
        got_row = got['detail_map'][key]
        for field in ['Units_On_Hand', 'Daily_Rate_Units_Per_Day', 'Current_Days_On_Hand', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand', 'Remaining_Demand_Units', 'Additional_Units_Needed', 'Loads_Required']:
            exp_val = exp_row.get(field)
            got_val = got_row.get(field)
            if exp_val == '' or exp_val is None:
                assert got_val in (None, '', ''), f'{field} should be blank for {key}'
            else:
                assert_approx(got_val, to_number(exp_val), 1e-4, f'{field} for {key}')
        for field in ['Projected_OOS_Date', 'Required_Delivery_Date']:
            exp_val = exp_row.get(field)
            got_val = got_row.get(field)
            assert to_iso(parse_date(got_val)) == to_iso(parse_date(exp_val)), f'{field} mismatch for {key}'
        for field in ['Earlier_Delivery_Required']:
            exp_val = exp_row.get(field)
            got_val = got_row.get(field)
            assert to_bool(got_val) == to_bool(exp_val), f'{field} mismatch for {key}'

    # Action sheet checks
    assert set(got['action_map'].keys()) == expected['action_keys'], 'Action key set mismatch'
    for key in expected['action_keys']:
        exp_act = expected['action_map'][key]
        got_act = got['action_map'].get(key)
        assert got_act is not None, f'Missing action row: {key}'
        for field in ['Units_On_Hand', 'Daily_Rate_Units_Per_Day', 'Current_Days_On_Hand', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand', 'Remaining_Demand_Units', 'Additional_Units_Needed', 'Loads_Required']:
            if field in got_act:
                assert_approx(got_act[field], to_number(exp_act.get(field)), 1e-4, f'Action {field} for {key}')
        for field in ['Projected_OOS_Date', 'Required_Delivery_Date']:
            if field in got_act:
                assert to_iso(parse_date(got_act[field])) == to_iso(parse_date(exp_act.get(field))), f'Action {field} mismatch for {key}'
        for field in ['Earlier_Delivery_Required']:
            if field in got_act:
                assert to_bool(got_act[field]) == to_bool(exp_act.get(field)), f'Action {field} mismatch for {key}'

    print('All checks passed.')


if __name__ == '__main__':
    main()
