#!/usr/bin/env python3
"""Solver for frozen_meal_tracker_refresh."""
import math
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook

EPS = 1e-9


def normalize_text(value: Any) -> str:
    return str(value or '').strip().upper()


def to_number(value: Any) -> float:
    if value in (None, ''):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(',', '').strip()
    return float(text) if text else 0.0


def parse_date(value: Any) -> date | None:
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def to_iso(value: date | None) -> str:
    return value.isoformat() if value else ''


def diff_days(start: date, end: date) -> int:
    return (end - start).days


def add_days(start: date, days: int) -> date:
    return start + timedelta(days=days)


def round_to(value: float, decimals: int = 4) -> float:
    factor = 10 ** decimals
    return round(value * factor) / factor


def main():
    input_root = Path(sys.argv[1])
    output_root = Path(sys.argv[2])

    wbt = load_workbook(input_root / 'Frozen_Meal_Template.xlsx')
    guide_ws = wbt['Pallet Guide']
    wbs = load_workbook(input_root / 'Frozen_Meal_Current_Stock.xlsx', data_only=True)
    stock_ws = wbs['Stock']
    wbl = load_workbook(input_root / 'Frozen_Meal_Recovery_Log.xlsx', data_only=True)
    log_ws = wbl['Recovery Log']

    as_of = parse_date(stock_ws['B1'].value)
    horizon = parse_date(stock_ws['D1'].value)
    planning_days = diff_days(as_of, horizon)

    pallet_by_sku: Dict[str, float] = {}
    for r in range(2, guide_ws.max_row + 1):
        sku = normalize_text(guide_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        cpp = to_number(guide_ws.cell(row=r, column=2).value)
        pallet_by_sku[sku] = cpp

    inventory = []
    for r in range(4, stock_ws.max_row + 1):
        sku = normalize_text(stock_ws.cell(row=r, column=1).value)
        if not sku:
            continue
        units = to_number(stock_ws.cell(row=r, column=2).value)
        rate = to_number(stock_ws.cell(row=r, column=3).value)
        inventory.append((sku, units, rate))

    raw_logs = []
    for r in range(2, log_ws.max_row + 1):
        lid = normalize_text(log_ws.cell(row=r, column=1).value)
        if not lid:
            continue
        rev = to_number(log_ws.cell(row=r, column=2).value)
        sku = normalize_text(log_ws.cell(row=r, column=3).value)
        d = parse_date(log_ws.cell(row=r, column=4).value)
        units = to_number(log_ws.cell(row=r, column=5).value)
        stage = normalize_text(log_ws.cell(row=r, column=6).value)
        raw_logs.append((lid, rev, sku, d, units, stage))

    by_lid: Dict[str, Tuple[float, str, date | None, float, str]] = {}
    for lid, rev, sku, d, units, stage in raw_logs:
        if d is None:
            continue
        prev = by_lid.get(lid)
        if prev is None or rev > prev[0]:
            by_lid[lid] = (rev, sku, d, units, stage)

    valid_stages = {'BOOKED', 'LOADED'}
    inbound_by_sku: Dict[str, List[Tuple[date, float]]] = {}
    for lid, (rev, sku, d, units, stage) in by_lid.items():
        if stage not in valid_stages:
            continue
        if sku not in inbound_by_sku:
            inbound_by_sku[sku] = []
        inbound_by_sku[sku].append((d, units))

    detail_rows = []
    action_rows = []

    for sku, units, rate in inventory:
        inbound = inbound_by_sku.get(sku, [])
        inbound_in_horizon = sum(u for d, u in inbound if d <= horizon)
        earliest = min((d for d, u in inbound if d <= horizon), default=None)
        current_doh = units / rate if rate > 0 else None
        proj_oos = add_days(as_of, int(current_doh)) if current_doh is not None else None
        delivered_doh = (units + inbound_in_horizon) / rate if rate > 0 else None
        remaining_demand = rate * planning_days
        additional = max(0.0, remaining_demand - units - inbound_in_horizon) if rate > 0 else 0.0
        cpp = pallet_by_sku.get(sku, 48)
        loads = math.ceil(additional / cpp) if additional > 0 else 0
        req_date = proj_oos if loads > 0 else None
        earlier = loads > 0 and (earliest is None or (req_date and req_date < earliest))

        detail_rows.append([
            sku,
            units,
            rate,
            round_to(current_doh) if current_doh is not None else '',
            to_iso(proj_oos),
            inbound_in_horizon,
            round_to(delivered_doh) if delivered_doh is not None else '',
            round_to(remaining_demand),
            round_to(additional),
            loads,
            to_iso(req_date),
            earlier,
        ])
        if loads > 0:
            action_rows.append([sku, to_iso(req_date), loads, round_to(additional), earlier])

    # Clear and repopulate Coverage_Detail
    cov_ws = wbt['Coverage_Detail']
    for r in range(1, cov_ws.max_row + 1):
        for c in range(1, cov_ws.max_column + 1):
            cov_ws.cell(row=r, column=c, value=None)

    cov_ws['A1'] = 'Field'
    cov_ws['B1'] = 'Value'
    cov_ws['A2'] = 'AsOfDate'
    cov_ws['B2'] = to_iso(as_of)
    cov_ws['A3'] = 'HorizonEnd'
    cov_ws['B3'] = to_iso(horizon)
    cov_ws['A4'] = 'PlanningDays'
    cov_ws['B4'] = planning_days

    headers = [
        'SKU', 'Units_On_Hand', 'Daily_Rate_Units_Per_Day', 'Current_Days_On_Hand',
        'Projected_OOS_Date', 'Inbound_Units_By_Horizon', 'Delivered_Days_On_Hand',
        'Remaining_Demand_Units', 'Additional_Units_Needed', 'Loads_Required',
        'Required_Delivery_Date', 'Earlier_Delivery_Required',
    ]
    for i, h in enumerate(headers, 1):
        cov_ws.cell(row=6, column=i, value=h)

    for ridx, row in enumerate(detail_rows, 7):
        for cidx, val in enumerate(row, 1):
            cov_ws.cell(row=ridx, column=cidx, value=val)

    # Clear and repopulate Recovery_Loads
    rec_ws = wbt['Recovery_Loads']
    for r in range(1, rec_ws.max_row + 1):
        for c in range(1, rec_ws.max_column + 1):
            rec_ws.cell(row=r, column=c, value=None)

    rec_headers = ['SKU', 'Required_Delivery_Date', 'Loads_Required', 'Additional_Units_Needed', 'Earlier_Delivery_Required']
    for i, h in enumerate(rec_headers, 1):
        rec_ws.cell(row=1, column=i, value=h)
    for ridx, row in enumerate(action_rows, 2):
        for cidx, val in enumerate(row, 1):
            rec_ws.cell(row=ridx, column=cidx, value=val)

    output_file = output_root / 'frozen_meal_recovery_tracker.xlsx'
    wbt.save(output_file)

    print('Wrote', output_file)


if __name__ == '__main__':
    main()
