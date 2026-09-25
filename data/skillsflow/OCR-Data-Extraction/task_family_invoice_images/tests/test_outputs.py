"""
Test verification for invoice extraction task.
"""

import os
from typing import Any
from openpyxl import load_workbook


def rows_match(actual_rows, expected_rows, tol=0.01):
    """Compare rows with numeric tolerance for amount fields."""
    if len(actual_rows) != len(expected_rows):
        return False, f"Row count mismatch: {len(actual_rows)} vs {len(expected_rows)}"
    for i, (ar, er) in enumerate(zip(actual_rows, expected_rows)):
        if len(ar) != len(er):
            return False, f"Row {i} column count mismatch: {len(ar)} vs {len(er)}"
        for j, (a, e) in enumerate(zip(ar, er)):
            a_s, e_s = str(a).strip(), str(e).strip()
            # Try numeric comparison
            try:
                a_f, e_f = float(a_s.replace(',', '')), float(e_s.replace(',', ''))
                if abs(a_f - e_f) > tol:
                    return False, f"Row {i} col {j}: numeric mismatch {a_s} vs {e_s}"
            except (ValueError, TypeError):
                if a_s != e_s:
                    return False, f"Row {i} col {j}: text mismatch '{a_s}' vs '{e_s}'"
    return True, ""


def _cell_to_string(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _read_sheet_as_string_rows(path: str, sheet_name: str, *, data_only: bool = True) -> list[list[str]]:
    wb = load_workbook(path, data_only=data_only)
    try:
        ws = wb[sheet_name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        rows: list[list[str]] = []
        for r in range(1, max_row + 1):
            rows.append([_cell_to_string(ws.cell(row=r, column=c).value) for c in range(1, max_col + 1)])
        return rows
    finally:
        wb.close()


def _get_sheetnames(path: str, *, data_only: bool = True) -> list[str]:
    wb = load_workbook(path, data_only=data_only)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def test_outputs():
    """Test that the invoice extraction outputs are correct."""
    
    OUTPUT_FILE = "/app/workspace/invoice_summary.xlsx"
    EXPECTED_FILE = os.path.join(os.path.dirname(__file__), "invoice_oracle.xlsx")
    
    # Requirement: output file exists
    assert os.path.exists(OUTPUT_FILE), "invoice_summary.xlsx not found at /app/workspace"
    
    # Requirement: single sheet named "invoices"
    actual_sheets = _get_sheetnames(OUTPUT_FILE)
    assert actual_sheets == ["invoices"], (
        f"Requirement failed: workbook must contain exactly one sheet named 'invoices'.\n"
        f"Actual sheets: {actual_sheets}"
    )
    
    # Read both files
    actual_rows = _read_sheet_as_string_rows(OUTPUT_FILE, "invoices")
    expected_rows = _read_sheet_as_string_rows(EXPECTED_FILE, "invoices")
    
    # Requirement: correct header
    expected_header = ["filename", "date", "total_amount"]
    actual_header = actual_rows[0] if actual_rows else []
    assert actual_header == expected_header, (
        f"Requirement failed: header/schema mismatch.\n"
        f"Actual header:   {actual_header}\n"
        f"Expected header: {expected_header}"
    )
    
    # Requirement: rows ordered by filename
    data_rows = actual_rows[1:]
    data_rows = [r for r in data_rows if len(r) >= 1 and r[0].strip() != ""]
    filenames = [r[0] for r in data_rows]
    assert filenames == sorted(filenames), (
        f"Requirement failed: rows are not ordered by filename.\n"
        f"Actual order:   {filenames}\n"
        f"Sorted order:   {sorted(filenames)}"
    )
    
    # Requirement: match with oracle (with numeric tolerance)
    ok, msg = rows_match(actual_rows, expected_rows)
    assert ok, f"Oracle mismatch: {msg}"
