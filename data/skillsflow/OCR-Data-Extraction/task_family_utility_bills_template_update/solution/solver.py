from __future__ import annotations

import csv
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR / 'tools'))
from ocr_utils import as_two_decimal_string, extract_amount_by_keywords, find_best_date, list_images, ocr_extract_text


DATA_DIR = Path('/app/workspace/dataset/img')
TEMPLATE_PATH = Path('/app/workspace/bill_tracker_template.xlsx')
OUTPUT_PATH = Path('/app/workspace/bill_tracker_filled.xlsx')

DATE_PATTERNS = [
    (r'STATEMENT\s*DATE[:\s]*([0-3]?\d[/\-][01]?\d[/\-]\d{2,4})', 50, True),
    (r'BILL\s*DATE[:\s]*([0-3]?\d[/\-][01]?\d[/\-]\d{2,4}|\d{4}-\d{2}-\d{2})', 50, True),
    (r'DATE[:\s]*([0-3]?\d[/\-][01]?\d[/\-]\d{2,4}|\d{4}-\d{2}-\d{2})', 10, True),
    (r'([0-3]?\d/[01]?\d/20\d{2})', 5, True),
    (r'([0-3]?\d-[01]?\d-20\d{2})', 5, True),
    (r'(20\d{2}-\d{2}-\d{2})', 5, True),
    (r'([01]?\d/[0-3]?\d/20\d{2})', 4, False),
]
AMOUNT_KEYWORDS = [r'PAY\s+THIS\s+AMOUNT', r'AMOUNT\s+DUE', r'TOTAL\s+DUE', r'CURRENT\s+CHARGES']
EXCLUDE_KEYWORDS = [r'PREVIOUS\s+BALANCE', r'LATE\s+FEE', r'TAX', r'CREDIT']


def extract_row(image_path: Path):
    text = ocr_extract_text(str(image_path))
    dt = find_best_date(text, DATE_PATTERNS)
    amount = extract_amount_by_keywords(text, AMOUNT_KEYWORDS, EXCLUDE_KEYWORDS)
    return [
        image_path.name,
        dt.strftime('%Y-%m-%d') if dt else None,
        as_two_decimal_string(amount) if amount is not None else None,
    ]


def main() -> None:
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb['bills']
    if ws.max_row and ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for image_path in list_images(DATA_DIR, recursive=False):
        ws.append(extract_row(image_path))
    wb.save(OUTPUT_PATH)


if __name__ == '__main__':
    main()
