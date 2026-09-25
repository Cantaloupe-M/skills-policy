from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR / 'tools'))
from ocr_utils import as_two_decimal_string, extract_amount_by_keywords, find_best_date, list_images, ocr_extract_text


DATA_DIR = Path('/app/workspace/dataset/cases')
OUTPUT_PATH = Path('/app/workspace/case_settlement.xlsx')

DATE_PATTERNS = [
    (r'ISSUE\s*DATE[:\s]*([0-3]?\d[/\-][01]?\d[/\-]\d{2,4}|\d{4}-\d{2}-\d{2})', 50, True),
    (r'DATE[:\s]*([0-3]?\d[/\-][01]?\d[/\-]\d{2,4}|\d{4}-\d{2}-\d{2})', 20, True),
    (r'\b(20\d{2}-\d{2}-\d{2})\b', 5, True),
    (r'\b([0-3]?\d-[01]?\d-20\d{2})\b', 5, True),
    (r'\b([01]?\d/[0-3]?\d/20\d{2})\b', 4, False),
]
PURCHASE_KEYWORDS = [r'PURCHASE\s+RECEIPT', r'TAX\s+INVOICE', r'STORE\s+RECEIPT']
CREDIT_KEYWORDS = [r'CREDIT\s+NOTE', r'REFUND\s+ADJUSTMENT', r'CREDIT\s+MEMO']
PURCHASE_AMOUNT_KEYWORDS = [r'GRAND\s+TOTAL', r'TOTAL\s+DUE', r'AMOUNT\s+DUE']
CREDIT_AMOUNT_KEYWORDS = [r'CREDIT\s+AMOUNT', r'REFUND\s+TOTAL', r'TOTAL\s+CREDIT']
COMMON_EXCLUDES = [r'SUBTOTAL', r'TAX', r'DISCOUNT']
REF_PATTERNS = [
    r'CREDIT\s*NO(?:\s*:\s*|\s+)([A-Z0-9\-]+)',
    r'RECEIPT\s*NO(?:\s*:\s*|\s+)([A-Z0-9\-]+)',
    r'REFERENCE(?:\s*:\s*|\s+)([A-Z0-9\-]+)',
]


def normalize_ref(value: str | None) -> str | None:
    if not value:
        return value
    value = re.sub(r'[^A-Z0-9\-]', '', value.upper())
    parts = value.split('-')
    normalized = []
    for idx, part in enumerate(parts):
        if idx >= 2 or (idx == 1 and any(ch.isdigit() for ch in part)):
            part = part.replace('O', '0').replace('I', '1').replace('L', '1')
        normalized.append(part)
    if len(normalized) >= 3:
        suffix = re.sub(r'\D', '', normalized[-1].replace('O', '0').replace('I', '1').replace('L', '1'))
        if suffix:
            normalized[-1] = suffix[-3:].zfill(3)
    return '-'.join(normalized)


def extract_document_ref(text: str):
    for pattern in REF_PATTERNS:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            candidate = normalize_ref(match.group(1))
            if candidate and len(candidate) >= 4:
                return candidate
    return None


def classify_document(text: str):
    if any(re.search(pattern, text, re.IGNORECASE) for pattern in CREDIT_KEYWORDS):
        return 'credit'
    if any(re.search(pattern, text, re.IGNORECASE) for pattern in PURCHASE_KEYWORDS):
        return 'purchase'
    return None


def main() -> None:
    seen_refs = set()
    events = []
    aggregates = defaultdict(lambda: {'purchase': Decimal('0.00'), 'credit': Decimal('0.00'), 'dates': []})
    for image_path in list_images(DATA_DIR, recursive=True):
        rel_path = image_path.relative_to(DATA_DIR).as_posix()
        case_id = rel_path.split('/')[0]
        text = ocr_extract_text(str(image_path))
        document_type = classify_document(text)
        if not document_type:
            continue
        document_ref = extract_document_ref(text)
        if document_ref and document_ref in seen_refs:
            continue
        if document_ref:
            seen_refs.add(document_ref)
        dt = find_best_date(text, DATE_PATTERNS)
        amount_keywords = PURCHASE_AMOUNT_KEYWORDS if document_type == 'purchase' else CREDIT_AMOUNT_KEYWORDS
        amount = extract_amount_by_keywords(text, amount_keywords, COMMON_EXCLUDES)
        events.append([
            case_id,
            rel_path,
            document_type,
            document_ref,
            dt.strftime('%Y-%m-%d') if dt else None,
            as_two_decimal_string(amount) if amount is not None else None,
        ])
        if amount is not None:
            aggregates[case_id][document_type] += amount
        if dt is not None:
            aggregates[case_id]['dates'].append(dt)
    events.sort(key=lambda row: (row[0], row[1]))
    wb = Workbook()
    ws_events = wb.active
    ws_events.title = 'events'
    ws_events.append(['case_id', 'relative_path', 'document_type', 'document_ref', 'date', 'amount'])
    for row in events:
        ws_events.append(row)
    ws_summary = wb.create_sheet('net_summary')
    ws_summary.append(['case_id', 'purchase_total', 'credit_total', 'net_amount', 'latest_date'])
    for case_id in sorted(aggregates.keys()):
        purchase_total = aggregates[case_id]['purchase']
        credit_total = aggregates[case_id]['credit']
        net_amount = purchase_total - credit_total
        dates = aggregates[case_id]['dates']
        latest_date = max(dates).strftime('%Y-%m-%d') if dates else None
        ws_summary.append([
            case_id,
            as_two_decimal_string(purchase_total),
            as_two_decimal_string(credit_total),
            as_two_decimal_string(net_amount),
            latest_date,
        ])
    wb.save(OUTPUT_PATH)


if __name__ == '__main__':
    main()
