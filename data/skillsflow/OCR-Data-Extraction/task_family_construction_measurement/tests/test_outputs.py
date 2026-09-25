"""
Test verification for construction measurement extraction task.
"""

import os
from typing import Any
from openpyxl import load_workbook


def rows_match(actual_rows, expected_rows, tol=0.01):
    """Compare rows with numeric tolerance for amount fields."""
    if len(actual_rows) != len(expected_rows):
        return False, f"Row count mismatch: {len(actual_rows)} vs {len(expected_rows)}"
    for i, (ar, er) in enumerate(zip(actual_rows, expected_rows)):
        if len(ar) != len(er):
            return False, f"Row {i} column count mismatch: {len(ar)} vs {len(er)}"
        for j, (a, e) in enumerate(zip(ar, er)):
            a_s, e_s = str(a).strip(), str(e).strip()
            # Try numeric comparison
            try:
                a_f, e_f = float(a_s.replace(',', '')), float(e_s.replace(',', ''))
                if abs(a_f - e_f) > tol:
                    return False, f"Row {i} col {j}: numeric mismatch {a_s} vs {e_s}"
            except (ValueError, TypeError):
                if a_s != e_s:
                    return False, f"Row {i} col {j}: text mismatch '{a_s}' vs '{e_s}'"
    return True, ""


def _cell_to_string(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _read_sheet_as_string_rows(path: str, sheet_name: str, *, data_only: bool = True) -> list[list[str]]:
    wb = load_workbook(path, data_only=data_only)
    try:
        ws = wb[sheet_name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        rows: list[list[str]] = []
        for r in range(1, max_row + 1):
            rows.append([_cell_to_string(ws.cell(row=r, column=c).value) for c in range(1, max_col + 1)])
        return rows
    finally:
        wb.close()


def _get_sheetnames(path: str, *, data_only: bool = True) -> list[str]:
    wb = load_workbook(path, data_only=data_only)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def test_outputs():
    """Test that the construction measurement extraction outputs are correct."""
    
    OUTPUT_FILE = "/app/workspace/construction_summary.xlsx"
    EXPECTED_FILE = os.path.join(os.path.dirname(__file__), "construction_oracle.xlsx")
    
    # Requirement: output file exists
    assert os.path.exists(OUTPUT_FILE), "construction_summary.xlsx not found at /app/workspace"
    
    # Requirement: two sheets named "details" and "summary"
    actual_sheets = _get_sheetnames(OUTPUT_FILE)
    assert actual_sheets == ["details", "summary"], (
        f"Requirement failed: workbook must contain exactly two sheets named 'details' and 'summary'.\n"
        f"Actual sheets: {actual_sheets}"
    )
    
    # Test details sheet
    details_rows = _read_sheet_as_string_rows(OUTPUT_FILE, "details")
    expected_details = _read_sheet_as_string_rows(EXPECTED_FILE, "details")
    
    # Requirement: correct header for details
    expected_header = ["filename", "project_code", "item_description", "quantity", "unit_price"]
    actual_header = details_rows[0] if details_rows else []
    assert actual_header == expected_header, (
        f"Requirement failed: details header/schema mismatch.\n"
        f"Actual header:   {actual_header}\n"
        f"Expected header: {expected_header}"
    )
    
    # Test summary sheet
    summary_rows = _read_sheet_as_string_rows(OUTPUT_FILE, "summary")
    expected_summary = _read_sheet_as_string_rows(EXPECTED_FILE, "summary")
    
    # Requirement: correct header for summary
    expected_summary_header = ["project_code", "date", "total_amount"]
    actual_summary_header = summary_rows[0] if summary_rows else []
    assert actual_summary_header == expected_summary_header, (
        f"Requirement failed: summary header/schema mismatch.\n"
        f"Actual header:   {actual_summary_header}\n"
        f"Expected header: {expected_summary_header}"
    )
    
    # Requirement: summary rows ordered by project_code
    summary_data = summary_rows[1:]
    summary_data = [r for r in summary_data if len(r) >= 1 and r[0].strip() != ""]
    project_codes = [r[0] for r in summary_data]
    assert project_codes == sorted(project_codes), (
        f"Requirement failed: summary rows are not ordered by project_code.\n"
        f"Actual order:   {project_codes}\n"
        f"Sorted order:   {sorted(project_codes)}"
    )
    
    # Requirement: match with oracle for both sheets (with numeric tolerance)
    ok, msg = rows_match(details_rows, expected_details)
    assert ok, f"Details oracle mismatch: {msg}"

    ok, msg = rows_match(summary_rows, expected_summary)
    assert ok, f"Summary oracle mismatch: {msg}"
