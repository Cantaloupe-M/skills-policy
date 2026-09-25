import os
import tempfile
import zipfile

from openpyxl import load_workbook

RESULT_FILE = "/root/results.pptx"
INPUT_FILE = "/root/input.pptx"
TARGET_EMBEDDING = 'Microsoft_Excel_Worksheet.xlsx'
SHEET_NAME = 'Region Factors'
TARGET_TITLE = 'Approved Adjustment Matrix'
ARCHIVE_TITLE = 'Archived Adjustment Matrix'
FROM_TOKEN = 'North Coast Zone'
TO_TOKEN = 'Harbor Fringe Zone'
EXPECTED_RATE = 1.24


def normalize(text):
    return ''.join(ch.lower() for ch in str(text) if ch.isalnum())


def token_matches(value, token):
    return value is not None and normalize(token) in normalize(value)


def workbook_from_pptx(pptx_path, data_only):
    with zipfile.ZipFile(pptx_path, 'r') as zf:
        payload = zf.read(f'ppt/embeddings/{TARGET_EMBEDDING}')
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp.write(payload)
        temp_path = tmp.name
    wb = load_workbook(temp_path, data_only=data_only)
    return wb, temp_path


def find_title(ws, title):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == title:
                return cell.row, cell.column
    raise AssertionError(f'Title {title} not found')


def matrix_bounds(ws, title):
    title_row, title_col = find_title(ws, title)
    anchor_row = title_row + 1
    anchor_col = title_col
    row_labels = []
    row = anchor_row + 1
    while ws.cell(row=row, column=anchor_col).value not in (None, ''):
        row_labels.append(row)
        row += 1
    col_labels = []
    col = anchor_col + 1
    while ws.cell(row=anchor_row, column=col).value not in (None, ''):
        col_labels.append(col)
        col += 1
    return anchor_row, anchor_col, row_labels, col_labels


def find_pair(ws, title, from_token, to_token):
    anchor_row, anchor_col, row_labels, col_labels = matrix_bounds(ws, title)
    from_row = next(row for row in row_labels if token_matches(ws.cell(row=row, column=anchor_col).value, from_token))
    to_row = next(row for row in row_labels if token_matches(ws.cell(row=row, column=anchor_col).value, to_token))
    from_col = next(col for col in col_labels if token_matches(ws.cell(row=anchor_row, column=col).value, from_token))
    to_col = next(col for col in col_labels if token_matches(ws.cell(row=anchor_row, column=col).value, to_token))
    return (from_row, to_col), (to_row, from_col), (anchor_row, anchor_col, row_labels, col_labels)


def is_formula(value):
    return isinstance(value, str) and value.startswith('=')


def assert_same(lhs, rhs, coord):
    if isinstance(lhs, (int, float)) and isinstance(rhs, (int, float)):
        assert abs(lhs - rhs) < 1e-6, f'Unexpected change at {coord}: {lhs} -> {rhs}'
    else:
        lhs_s = str(lhs).strip() if lhs is not None else None
        rhs_s = str(rhs).strip() if rhs is not None else None
        assert lhs_s == rhs_s, f'Unexpected change at {coord}: {lhs} -> {rhs}'


def test_output_exists():
    assert os.path.exists(RESULT_FILE), 'results.pptx was not created'
    assert zipfile.is_zipfile(RESULT_FILE), 'results.pptx is not a valid PPTX archive'


def test_live_matrix_updated_and_archive_matrix_unchanged():
    input_wb, input_path = workbook_from_pptx(INPUT_FILE, data_only=True)
    output_wb, output_path = workbook_from_pptx(RESULT_FILE, data_only=True)
    try:
        input_ws = input_wb[SHEET_NAME]
        output_ws = output_wb[SHEET_NAME]
        direct_coord, inverse_coord, target_bounds = find_pair(output_ws, TARGET_TITLE, FROM_TOKEN, TO_TOKEN)
        direct_value = output_ws.cell(*direct_coord).value
        inverse_value = output_ws.cell(*inverse_coord).value
        assert direct_value is not None, f'Direct rate cell is empty (no cached value)'
        assert abs(direct_value - EXPECTED_RATE) < 1e-4, f'Expected {EXPECTED_RATE}, got {direct_value}'
        # 双通路验证：先尝试缓存值，若为 None 则回退到公式语义检查
        if inverse_value is not None:
            assert abs(inverse_value - (1.0 / EXPECTED_RATE)) < 1e-3, f'Expected {1.0 / EXPECTED_RATE}, got {inverse_value}'
        else:
            # 缓存值为空，回退到公式语义检查
            output_wb.close()
            os.unlink(output_path)
            output_wb, output_path = workbook_from_pptx(RESULT_FILE, data_only=False)
            output_ws = output_wb[SHEET_NAME]
            inverse_formula = output_ws.cell(*inverse_coord).value
            assert is_formula(inverse_formula), (
                f'Inverse rate cell is neither a cached value nor a formula: {inverse_formula}'
            )
            # 检查公式是否包含对直接值单元格的正确引用（1/Dxx 或类似结构）
            import re
            direct_cell = output_ws.cell(*direct_coord).coordinate
            norm_formula = str(inverse_formula).replace(' ', '').replace('$', '').upper()
            # 允许 1/直接单元格 或 ROUND(1/直接单元格) 等变体
            pattern = rf'(ROUND\()?1/{direct_cell}(,\d+)?\)?'
            assert re.search(pattern, norm_formula), (
                f'Inverse formula does not reference the direct rate cell correctly: {inverse_formula}'
            )

        _, _, archive_bounds = find_pair(output_ws, ARCHIVE_TITLE, FROM_TOKEN, TO_TOKEN)

        for row in archive_bounds[2]:
            for col in archive_bounds[3]:
                coord = output_ws.cell(row=row, column=col).coordinate
                assert_same(input_ws.cell(row=row, column=col).value, output_ws.cell(row=row, column=col).value, f'archive!{coord}')

        for row in target_bounds[2]:
            for col in target_bounds[3]:
                coord = (row, col)
                if coord in {direct_coord, inverse_coord}:
                    continue
                cell_ref = output_ws.cell(row=row, column=col).coordinate
                assert_same(input_ws.cell(row=row, column=col).value, output_ws.cell(row=row, column=col).value, f'live!{cell_ref}')
    finally:
        input_wb.close()
        output_wb.close()
        os.unlink(input_path)
        os.unlink(output_path)


def test_formula_cells_preserved():
    input_wb, input_path = workbook_from_pptx(INPUT_FILE, data_only=False)
    output_wb, output_path = workbook_from_pptx(RESULT_FILE, data_only=False)
    try:
        input_formulas = {}
        output_formulas = {}
        for sheet_name in input_wb.sheetnames:
            input_ws = input_wb[sheet_name]
            output_ws = output_wb[sheet_name]
            for row in range(1, input_ws.max_row + 1):
                for col in range(1, input_ws.max_column + 1):
                    coord = input_ws.cell(row=row, column=col).coordinate
                    input_value = input_ws.cell(row=row, column=col).value
                    output_value = output_ws.cell(row=row, column=col).value
                    if is_formula(input_value):
                        input_formulas[(sheet_name, coord)] = input_value
                    if is_formula(output_value):
                        output_formulas[(sheet_name, coord)] = output_value
        # Normalize formulas for comparison (case-insensitive, remove spaces and $)
        _input_norm = {k: str(v).replace(' ', '').replace('$', '').upper() for k, v in input_formulas.items()}
        _output_norm = {k: str(v).replace(' ', '').replace('$', '').upper() for k, v in output_formulas.items()}
        assert _input_norm == _output_norm, 'Formula cells changed or were replaced with values'
    finally:
        input_wb.close()
        output_wb.close()
        os.unlink(input_path)
        os.unlink(output_path)
