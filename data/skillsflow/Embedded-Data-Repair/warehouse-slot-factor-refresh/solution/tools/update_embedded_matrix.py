#!/usr/bin/env python3
import argparse
import csv
import json
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
IGNORE_MARKERS = ('archive', 'archived', 'superseded', 'draft', 'obsolete')
PREFER_MARKERS = ('final', 'approved', 'live', 'current')
UPDATE_PATTERN = re.compile(
    r'([A-Za-z0-9]+)\s+to\s+([A-Za-z0-9]+)\s*=\s*([0-9]+(?:\.[0-9]+)?)',
    re.IGNORECASE,
)


def normalize(text: object) -> str:
    return ''.join(ch.lower() for ch in str(text) if ch.isalnum())


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith('=')


def load_config(config_path: Path) -> dict:
    with open(config_path, 'r', encoding='utf-8') as handle:
        return json.load(handle)


def extract_text_chunks(work_dir: Path) -> list[str]:
    chunks: list[str] = []
    slide_dir = work_dir / 'ppt' / 'slides'
    for slide_xml in sorted(slide_dir.glob('slide*.xml')):
        tree = ET.parse(slide_xml)
        for node in tree.iterfind(f'.//{A_NS}t'):
            text = (node.text or '').strip()
            if not text:
                continue
            parts = [part.strip() for part in re.split(r'\|\||\n', text) if part.strip()]
            chunks.extend(parts)
    return chunks


def choose_update(chunks: list[str]) -> tuple[str, str, float]:
    candidates: list[tuple[int, int, str, str, float]] = []
    for index, chunk in enumerate(chunks):
        match = UPDATE_PATTERN.search(chunk)
        if not match:
            continue
        lower = chunk.lower()
        score = 0
        if any(marker in lower for marker in PREFER_MARKERS):
            score += 5
        if any(marker in lower for marker in IGNORE_MARKERS):
            score -= 5
        candidates.append((score, index, match.group(1), match.group(2), float(match.group(3))))
    if not candidates:
        raise ValueError('No usable update instruction found in slide text')
    candidates.sort(key=lambda item: (item[0], item[1]))
    _, _, from_token, to_token, rate = candidates[-1]
    return from_token, to_token, rate


def resolve_aliases(config: dict, from_token: str, to_token: str) -> tuple[str, str]:
    alias_path = config.get('alias_map')
    if not alias_path:
        return from_token, to_token
    alias_map = {}
    with open(alias_path, 'r', encoding='utf-8', newline='') as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            alias_map[normalize(row['note_token'])] = row['workbook_label']
    return alias_map.get(normalize(from_token), from_token), alias_map.get(normalize(to_token), to_token)


def select_embedding(work_dir: Path, config: dict) -> Path:
    embeddings_dir = work_dir / 'ppt' / 'embeddings'
    embeddings = sorted(embeddings_dir.glob('*.xlsx'))
    if not embeddings:
        raise ValueError('No embedded workbook found inside the PPTX')
    target_name = config.get('target_embedding')
    if target_name:
        target_path = embeddings_dir / target_name
        if not target_path.exists():
            raise ValueError(f'Target embedding {target_name} not found')
        return target_path
    if len(embeddings) != 1:
        raise ValueError('Multiple embedded workbooks found but no target embedding was configured')
    return embeddings[0]


def token_matches(value: object, token: str) -> bool:
    return value is not None and normalize(token) in normalize(value)


def matrix_labels_at_anchor(ws, anchor_row: int, anchor_col: int):
    row_labels = {}
    row = anchor_row + 1
    while row <= ws.max_row and ws.cell(row=row, column=anchor_col).value not in (None, ''):
        row_labels[row] = ws.cell(row=row, column=anchor_col).value
        row += 1
    col_labels = {}
    col = anchor_col + 1
    while col <= ws.max_column and ws.cell(row=anchor_row, column=col).value not in (None, ''):
        col_labels[col] = ws.cell(row=anchor_row, column=col).value
        col += 1
    return row_labels, col_labels


def locate_matrix_at_anchor(ws, anchor_row: int, anchor_col: int, from_token: str, to_token: str):
    row_labels, col_labels = matrix_labels_at_anchor(ws, anchor_row, anchor_col)
    from_rows = [row for row, value in row_labels.items() if token_matches(value, from_token)]
    to_rows = [row for row, value in row_labels.items() if token_matches(value, to_token)]
    from_cols = [col for col, value in col_labels.items() if token_matches(value, from_token)]
    to_cols = [col for col, value in col_labels.items() if token_matches(value, to_token)]
    if not from_rows or not to_rows or not from_cols or not to_cols:
        raise ValueError(f'Could not find tokens {from_token} / {to_token} at anchor {anchor_row},{anchor_col}')
    return {
        'from_row': from_rows[0],
        'to_row': to_rows[0],
        'from_col': from_cols[0],
        'to_col': to_cols[0],
    }


def find_title_cell(ws, title: str):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == title:
                return cell.row, cell.column
    raise ValueError(f'Could not find matrix title {title!r}')


def locate_matrix(ws, from_token: str, to_token: str):
    for anchor_row in range(1, ws.max_row + 1):
        for anchor_col in range(1, ws.max_column + 1):
            if ws.cell(row=anchor_row, column=anchor_col).value in (None, ''):
                continue
            try:
                matrix = locate_matrix_at_anchor(ws, anchor_row, anchor_col, from_token, to_token)
                direct = ws.cell(row=matrix['from_row'], column=matrix['to_col']).value
                inverse = ws.cell(row=matrix['to_row'], column=matrix['from_col']).value
                if direct is not None and inverse is not None:
                    return matrix
            except Exception:
                continue
    raise ValueError(f'Could not find a reciprocal matrix for {from_token} and {to_token}')


def locate_target(ws, config: dict, from_token: str, to_token: str):
    if config.get('matrix_title'):
        title_row, title_col = find_title_cell(ws, config['matrix_title'])
        return locate_matrix_at_anchor(ws, title_row + 1, title_col, from_token, to_token)
    if config.get('anchor_row') and config.get('anchor_col'):
        return locate_matrix_at_anchor(ws, int(config['anchor_row']), int(config['anchor_col']), from_token, to_token)
    return locate_matrix(ws, from_token, to_token)


def recalc_excel(recalc_script: Path, workbook_path: Path) -> None:
    result = subprocess.run(
        ['python3', str(recalc_script), str(workbook_path), '90'],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(f'Recalculation failed: {result.stderr or result.stdout}')
    payload = result.stdout.strip()
    if payload:
        data = json.loads(payload)
        if 'error' in data:
            raise RuntimeError(data['error'])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--config', required=True)
    args = parser.parse_args()

    base_dir = Path(__file__).resolve().parent
    config = load_config(Path(args.config))
    unpack_script = base_dir / 'unpack.py'
    pack_script = base_dir / 'pack.py'
    recalc_script = base_dir / 'recalc.py'

    work_dir = Path(tempfile.mkdtemp(prefix='embedded-matrix-'))
    try:
        subprocess.run(['python3', str(unpack_script), args.input, str(work_dir)], check=True)
        text_chunks = extract_text_chunks(work_dir)
        from_token, to_token, new_rate = choose_update(text_chunks)
        from_token, to_token = resolve_aliases(config, from_token, to_token)
        workbook_path = select_embedding(work_dir, config)

        wb = load_workbook(workbook_path)
        try:
            if config.get('target_sheet'):
                ws = wb[config['target_sheet']]
            else:
                ws = wb.active
            matrix = locate_target(ws, config, from_token, to_token)
            direct_cell = ws.cell(row=matrix['from_row'], column=matrix['to_col'])
            inverse_cell = ws.cell(row=matrix['to_row'], column=matrix['from_col'])
            if not is_formula(direct_cell.value):
                direct_cell.value = new_rate
            elif not is_formula(inverse_cell.value):
                inverse_cell.value = 1.0 / new_rate
            else:
                raise ValueError('Both candidate cells are formulas; no writable input cell found')
            wb.save(workbook_path)
        finally:
            wb.close()

        recalc_excel(recalc_script, workbook_path)
        subprocess.run(['python3', str(pack_script), str(work_dir), args.output, '--force'], check=True)
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


if __name__ == '__main__':
    main()
