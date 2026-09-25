import os
import tempfile
import zipfile

from openpyxl import load_workbook

RESULT_FILE = "/root/results.pptx"
INPUT_FILE = "/root/input.pptx"
TARGET_EMBEDDING = 'Live_Rebate_Bands.xlsx'
SHEET_NAME = 'Approved Grid'
HEADER_ROW = 3
ROW_LABEL_COL = 2
FROM_TOKEN = 'Gold'
TO_TOKEN = 'Platinum'
EXPECTED_RATE = 1.31
UNCHANGED_SHEETS = []
ARCHIVE_EMBEDDINGS = ['Archive_Rebate_Bands.xlsx']


def normalize(text):
    return ''.join(ch.lower() for ch in str(text) if ch.isalnum())


def token_matches(value, token):
    return value is not None and normalize(token) in normalize(value)


def extract_embedding_bytes(pptx_path, embedding_name):
    with zipfile.ZipFile(pptx_path, 'r') as zf:
        return zf.read(f'ppt/embeddings/{embedding_name}')


def workbook_from_embedding(pptx_path, embedding_name, data_only):
    payload = extract_embedding_bytes(pptx_path, embedding_name)
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp.write(payload)
        temp_path = tmp.name
    wb = load_workbook(temp_path, data_only=data_only)
    return wb, temp_path


def find_pair(ws, from_token, to_token):
    row_map = {
        row: ws.cell(row=row, column=ROW_LABEL_COL).value
        for row in range(HEADER_ROW + 1, ws.max_row + 1)
        if ws.cell(row=row, column=ROW_LABEL_COL).value not in (None, '')
    }
    col_map = {
        col: ws.cell(row=HEADER_ROW, column=col).value
        for col in range(ROW_LABEL_COL + 1, ws.max_column + 1)
        if ws.cell(row=HEADER_ROW, column=col).value not in (None, '')
    }
    from_row = next(row for row, value in row_map.items() if token_matches(value, from_token))
    to_row = next(row for row, value in row_map.items() if token_matches(value, to_token))
    from_col = next(col for col, value in col_map.items() if token_matches(value, from_token))
    to_col = next(col for col, value in col_map.items() if token_matches(value, to_token))
    return (from_row, to_col), (to_row, from_col)


def all_matrix_values(ws):
    values = {}
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        for col in range(ROW_LABEL_COL + 1, ws.max_column + 1):
            values[(row, col)] = ws.cell(row=row, column=col).value
    return values


def is_formula(value):
    return isinstance(value, str) and value.startswith('=')


def assert_same(lhs, rhs, coord):
    if isinstance(lhs, (int, float)) and isinstance(rhs, (int, float)):
        assert abs(lhs - rhs) < 1e-6, f'Unexpected change at {coord}: {lhs} -> {rhs}'
    else:
        lhs_s = str(lhs).strip() if lhs is not None else None
        rhs_s = str(rhs).strip() if rhs is not None else None
        assert lhs_s == rhs_s, f'Unexpected change at {coord}: {lhs} -> {rhs}'


def test_output_exists_and_contains_target_embedding():
    assert os.path.exists(RESULT_FILE), 'results.pptx was not created'
    assert zipfile.is_zipfile(RESULT_FILE), 'results.pptx is not a valid PPTX archive'
    with zipfile.ZipFile(RESULT_FILE, 'r') as zf:
        assert f'ppt/embeddings/{TARGET_EMBEDDING}' in zf.namelist(), 'Target embedding is missing from the output'


def test_target_pair_updated():
    wb, temp_path = workbook_from_embedding(RESULT_FILE, TARGET_EMBEDDING, data_only=True)
    try:
        ws = wb[SHEET_NAME]
        direct_coord, inverse_coord = find_pair(ws, FROM_TOKEN, TO_TOKEN)
        direct_value = ws.cell(*direct_coord).value
        inverse_value = ws.cell(*inverse_coord).value
        assert direct_value is not None, f'Direct rate cell is empty (no cached value)'
        assert abs(direct_value - EXPECTED_RATE) < 1e-4, f'Expected {EXPECTED_RATE}, got {direct_value}'
        # 双通路验证：先尝试缓存值，若为 None 则回退到公式语义检查
        if inverse_value is not None:
            assert abs(inverse_value - (1.0 / EXPECTED_RATE)) < 1e-3, f'Expected {1.0 / EXPECTED_RATE}, got {inverse_value}'
        else:
            # 缓存值为空，回退到公式语义检查
            wb.close()
            os.unlink(temp_path)
            wb, temp_path = workbook_from_embedding(RESULT_FILE, TARGET_EMBEDDING, data_only=False)
            ws = wb[SHEET_NAME]
            inverse_formula = ws.cell(*inverse_coord).value
            assert is_formula(inverse_formula), (
                f'Inverse rate cell is neither a cached value nor a formula: {inverse_formula}'
            )
            # 检查公式是否包含对直接值单元格的正确引用（1/Dxx 或类似结构）
            import re
            direct_cell = ws.cell(*direct_coord).coordinate
            norm_formula = str(inverse_formula).replace(' ', '').replace('$', '').upper()
            # 允许 1/直接单元格 或 ROUND(1/直接单元格) 等变体
            pattern = rf'(ROUND\()?1/{direct_cell}(,\d+)?\)?'
            assert re.search(pattern, norm_formula), (
                f'Inverse formula does not reference the direct rate cell correctly: {inverse_formula}'
            )
    finally:
        wb.close()
        os.unlink(temp_path)


def test_formula_cells_preserved_in_target_embedding():
    input_wb, input_path = workbook_from_embedding(INPUT_FILE, TARGET_EMBEDDING, data_only=False)
    output_wb, output_path = workbook_from_embedding(RESULT_FILE, TARGET_EMBEDDING, data_only=False)
    try:
        input_formulas = {}
        output_formulas = {}
        for sheet_name in input_wb.sheetnames:
            input_ws = input_wb[sheet_name]
            output_ws = output_wb[sheet_name]
            for row in range(1, input_ws.max_row + 1):
                for col in range(1, input_ws.max_column + 1):
                    coord = input_ws.cell(row=row, column=col).coordinate
                    input_value = input_ws.cell(row=row, column=col).value
                    output_value = output_ws.cell(row=row, column=col).value
                    if is_formula(input_value):
                        input_formulas[(sheet_name, coord)] = input_value
                    if is_formula(output_value):
                        output_formulas[(sheet_name, coord)] = output_value
        # Normalize formulas for comparison (case-insensitive, remove spaces and $)
        _input_norm = {k: str(v).replace(' ', '').replace('$', '').upper() for k, v in input_formulas.items()}
        _output_norm = {k: str(v).replace(' ', '').replace('$', '').upper() for k, v in output_formulas.items()}
        assert _input_norm == _output_norm, 'Formula cells changed or were replaced with values'
    finally:
        input_wb.close()
        output_wb.close()
        os.unlink(input_path)
        os.unlink(output_path)


def test_other_values_and_archive_embeddings_unchanged():
    input_wb, input_path = workbook_from_embedding(INPUT_FILE, TARGET_EMBEDDING, data_only=True)
    output_wb, output_path = workbook_from_embedding(RESULT_FILE, TARGET_EMBEDDING, data_only=True)
    try:
        input_ws = input_wb[SHEET_NAME]
        output_ws = output_wb[SHEET_NAME]
        direct_coord, inverse_coord = find_pair(output_ws, FROM_TOKEN, TO_TOKEN)
        input_values = all_matrix_values(input_ws)
        output_values = all_matrix_values(output_ws)
        for coord, input_value in input_values.items():
            if coord in {direct_coord, inverse_coord}:
                continue
            assert_same(input_value, output_values[coord], coord)
        for sheet_name in UNCHANGED_SHEETS:
            in_sheet = input_wb[sheet_name]
            out_sheet = output_wb[sheet_name]
            max_row = max(in_sheet.max_row, out_sheet.max_row)
            max_col = max(in_sheet.max_column, out_sheet.max_column)
            for row in range(1, max_row + 1):
                for col in range(1, max_col + 1):
                    coord = in_sheet.cell(row=row, column=col).coordinate
                    assert_same(in_sheet.cell(row=row, column=col).value, out_sheet.cell(row=row, column=col).value, f'{sheet_name}!{coord}')
    finally:
        input_wb.close()
        output_wb.close()
        os.unlink(input_path)
        os.unlink(output_path)

    for embedding_name in ARCHIVE_EMBEDDINGS:
        assert extract_embedding_bytes(INPUT_FILE, embedding_name) == extract_embedding_bytes(RESULT_FILE, embedding_name), (
            f'Archive embedding {embedding_name} should be unchanged'
        )
