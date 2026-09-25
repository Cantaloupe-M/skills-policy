#!/usr/bin/env python3
import argparse
import json
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
IGNORE_MARKERS = ('archive', 'archived', 'superseded', 'draft', 'obsolete')
PREFER_MARKERS = ('final', 'approved', 'live', 'current')
UPDATE_PATTERN = re.compile(
    r'([A-Za-z0-9]+)\s+to\s+([A-Za-z0-9]+)\s*=\s*([0-9]+(?:\.[0-9]+)?)',
    re.IGNORECASE,
)


def normalize(text: object) -> str:
    return ''.join(ch.lower() for ch in str(text) if ch.isalnum())


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith('=')


def extract_text_chunks(work_dir: Path) -> list[str]:
    chunks: list[str] = []
    slide_dir = work_dir / 'ppt' / 'slides'
    for slide_xml in sorted(slide_dir.glob('slide*.xml')):
        tree = ET.parse(slide_xml)
        for node in tree.iterfind(f'.//{A_NS}t'):
            text = (node.text or '').strip()
            if not text:
                continue
            parts = [part.strip() for part in re.split(r'\|\||\n', text) if part.strip()]
            chunks.extend(parts)
    return chunks


def choose_update(chunks: list[str]) -> tuple[str, str, float]:
    candidates: list[tuple[int, int, str, str, float]] = []
    for index, chunk in enumerate(chunks):
        match = UPDATE_PATTERN.search(chunk)
        if not match:
            continue
        lower = chunk.lower()
        score = 0
        if any(marker in lower for marker in PREFER_MARKERS):
            score += 5
        if any(marker in lower for marker in IGNORE_MARKERS):
            score -= 5
        candidates.append((score, index, match.group(1), match.group(2), float(match.group(3))))
    if not candidates:
        raise ValueError('No usable update instruction found in slide text')
    candidates.sort(key=lambda item: (item[0], item[1]))
    _, _, from_token, to_token, rate = candidates[-1]
    return from_token, to_token, rate


def token_matches(value: object, token: str) -> bool:
    if value is None:
        return False
    return normalize(token) in normalize(value)


def locate_matrix(ws, from_token: str, to_token: str):
    for anchor_row in range(1, ws.max_row + 1):
        for anchor_col in range(1, ws.max_column + 1):
            row_labels = {
                row: ws.cell(row=row, column=anchor_col).value
                for row in range(anchor_row + 1, ws.max_row + 1)
                if ws.cell(row=row, column=anchor_col).value not in (None, '')
            }
            col_labels = {
                col: ws.cell(row=anchor_row, column=col).value
                for col in range(anchor_col + 1, ws.max_column + 1)
                if ws.cell(row=anchor_row, column=col).value not in (None, '')
            }
            if len(row_labels) < 2 or len(col_labels) < 2:
                continue

            from_rows = [row for row, value in row_labels.items() if token_matches(value, from_token)]
            to_rows = [row for row, value in row_labels.items() if token_matches(value, to_token)]
            from_cols = [col for col, value in col_labels.items() if token_matches(value, from_token)]
            to_cols = [col for col, value in col_labels.items() if token_matches(value, to_token)]
            if not from_rows or not to_rows or not from_cols or not to_cols:
                continue

            for from_row in from_rows:
                for to_row in to_rows:
                    for from_col in from_cols:
                        for to_col in to_cols:
                            direct = ws.cell(row=from_row, column=to_col).value
                            inverse = ws.cell(row=to_row, column=from_col).value
                            if direct is None or inverse is None:
                                continue
                            return {
                                'anchor_row': anchor_row,
                                'anchor_col': anchor_col,
                                'from_row': from_row,
                                'to_row': to_row,
                                'from_col': from_col,
                                'to_col': to_col,
                            }
    return None


def locate_workbook_target(workbook_path: Path, from_token: str, to_token: str):
    wb = load_workbook(workbook_path)
    try:
        for ws in wb.worksheets:
            matrix = locate_matrix(ws, from_token, to_token)
            if matrix:
                return wb, ws, matrix
    except Exception:
        wb.close()
        raise
    wb.close()
    raise ValueError(f'Could not find a reciprocal matrix for {from_token} and {to_token}')


def recalc_excel(recalc_script: Path, workbook_path: Path) -> None:
    result = subprocess.run(
        ['python3', str(recalc_script), str(workbook_path), '90'],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(f'Recalculation failed: {result.stderr or result.stdout}')
    payload = result.stdout.strip()
    if payload:
        data = json.loads(payload)
        if 'error' in data:
            raise RuntimeError(data['error'])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()

    base_dir = Path(__file__).resolve().parent
    unpack_script = base_dir / 'unpack.py'
    pack_script = base_dir / 'pack.py'
    recalc_script = base_dir / 'recalc.py'

    work_dir = Path(tempfile.mkdtemp(prefix='embedded-matrix-'))
    try:
        subprocess.run(['python3', str(unpack_script), args.input, str(work_dir)], check=True)
        text_chunks = extract_text_chunks(work_dir)
        from_token, to_token, new_rate = choose_update(text_chunks)

        embeddings = sorted((work_dir / 'ppt' / 'embeddings').glob('*.xlsx'))
        if not embeddings:
            raise ValueError('No embedded workbook found inside the PPTX')
        workbook_path = embeddings[0]

        wb, ws, matrix = locate_workbook_target(workbook_path, from_token, to_token)
        try:
            direct_cell = ws.cell(row=matrix['from_row'], column=matrix['to_col'])
            inverse_cell = ws.cell(row=matrix['to_row'], column=matrix['from_col'])
            if not is_formula(direct_cell.value):
                direct_cell.value = new_rate
            elif not is_formula(inverse_cell.value):
                inverse_cell.value = 1.0 / new_rate
            else:
                raise ValueError('Both candidate cells are formulas; no writable input cell found')
            wb.save(workbook_path)
        finally:
            wb.close()

        recalc_excel(recalc_script, workbook_path)
        subprocess.run(['python3', str(pack_script), str(work_dir), args.output, '--force'], check=True)
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


if __name__ == '__main__':
    main()
